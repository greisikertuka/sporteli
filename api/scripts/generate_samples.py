"""Generate the SYNTHETIC department exports of contract §1 into ``api/samples/``.

Deterministic (fixed seed, fixed file timestamps): running it twice produces identical bytes.
Only the standard library and openpyxl are used.

The files imitate real municipal exports so the ingest pipeline is exercised honestly: title
rows (merged, containing "SINTETIKE"), header on row 3, units in headers, blank and total rows
whose sums equal the data, Albanian date/number formats, admin-unit spelling variants and
clearly fake personal columns that the PII gate must drop.

None of these numbers are official figures. Population totals are deliberately not the
official census/civil-registry values.

Usage:  cd api && uv run python scripts/generate_samples.py
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import random
import re
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SEED = 20260926
OUT = Path(__file__).resolve().parents[1] / "samples"
AS_OF = dt.date(2026, 8, 31)
YEAR = 2026
MONTHS = [dt.date(YEAR, m, 1) for m in range(1, 9)]
MONTH_SQ = ["Janar", "Shkurt", "Mars", "Prill", "Maj", "Qershor", "Korrik", "Gusht"]
FIXED_DT = dt.datetime(2026, 8, 31, 18, 0, 0)
FIXED_ZIP_TIME = (2026, 8, 31, 18, 0, 0)

SYNTH_NOTE = "Të dhëna sintetike për demonstrim — nuk janë shifra zyrtare"

# --------------------------------------------------------------------------------------------
# Reference: administrative units and synthetic populations (NOT official figures)
# --------------------------------------------------------------------------------------------

ADMIN_UNITS = [
    # name, census_2023 (synthetic), civil_registry (synthetic)
    ("Elbasan", 83_650, 132_400),
    ("Bradashesh", 4_050, 8_300),
    ("Funarë", 1_150, 2_900),
    ("Gjergjan", 3_350, 7_100),
    ("Gjinar", 1_300, 3_300),
    ("Gracen", 1_650, 3_900),
    ("Labinot-Fushë", 4_750, 9_200),
    ("Labinot-Mal", 1_900, 4_700),
    ("Papër", 4_250, 8_100),
    ("Shirgjan", 6_100, 10_600),
    ("Shushicë", 5_300, 9_800),
    ("Tregan", 2_300, 5_100),
    ("Zavalinë", 1_650, 3_300),
]
CENSUS_TOTAL = 121_400
REGISTRY_TOTAL = 208_700

UNIT_VARIANTS = {
    "Elbasan": [("Elbasan (qyteti)", 0.05), ("ELBASAN", 0.03)],
    "Shirgjan": [("Shirgjani", 0.25)],
    "Tregan": [("Tregani", 0.25)],
    "Labinot-Fushë": [("Labinot Fushe", 0.2)],
    "Funarë": [("Funare", 0.3)],
    "Zavalinë": [("Zavalina", 0.2)],
}

FIRST_NAMES = [
    "Arben", "Blerta", "Dritan", "Elona", "Fatjon", "Gentiana", "Ilir", "Jonida", "Klodian",
    "Lindita", "Mirela", "Ndriçim", "Orjeta", "Petrit", "Rezarta", "Sokol", "Teuta", "Valbona",
    "Xhoni", "Zamira", "Ardit", "Besa", "Eriona", "Gëzim", "Hana", "Kristi", "Luan", "Marsela",
    "Nertil", "Oltiana",
]  # fmt: skip
INITIALS = "ABÇDEFGHIJKLMNOPRSTUVXZ"

# --------------------------------------------------------------------------------------------
# Formatting helpers
# --------------------------------------------------------------------------------------------


def fmt_date(d: dt.date) -> str:
    return d.strftime("%d.%m.%Y")


def fmt_sq(value: float, decimals: int = 0) -> str:
    """Albanian number format: thousands '.', decimal ','."""
    text = f"{value:,.{decimals}f}"
    return text.replace(",", "\u0000").replace(".", ",").replace("\u0000", ".")


def tenths(x: float) -> int:
    return int(round(x * 10))


def pick_weighted(rng: random.Random, items: list, weights: list[float]):
    return rng.choices(items, weights=weights, k=1)[0]


def unit_spelling(rng: random.Random, unit: str) -> str:
    for variant, p in UNIT_VARIANTS.get(unit, []):
        if rng.random() < p:
            return variant
    if rng.random() < 0.02:
        return unit + " "  # stray trailing space, as in hand-typed exports
    return unit


def fake_name(rng: random.Random) -> str:
    return f"{rng.choice(FIRST_NAMES)} {rng.choice(INITIALS)}."


def fake_phone(rng: random.Random) -> str:
    return f"+355 69 000 {rng.randint(0, 99):02d} {rng.randint(0, 99):02d}"


# --------------------------------------------------------------------------------------------
# XLSX helpers
# --------------------------------------------------------------------------------------------

THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
TITLE_FILL = PatternFill("solid", fgColor="FFF2CC")
TOTAL_FILL = PatternFill("solid", fgColor="EDEDED")


def write_titles(ws, lines: list[str], ncols: int) -> None:
    for i, text in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=text)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=ncols)
        cell = ws.cell(row=i, column=1)
        cell.font = Font(bold=i == 1, size=12 if i == 1 else 10, color="7F6000")
        cell.fill = TITLE_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")


def write_header(ws, row: int, headers: list[str]) -> None:
    for c, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=text)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_total(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True)
        cell.fill = TOTAL_FILL
        cell.border = BORDER


def set_widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def save_xlsx(wb: Workbook, path: Path) -> None:
    """Save with fixed metadata and zip timestamps so the bytes are reproducible."""
    wb.properties.creator = "Sportel · generate_samples.py (SINTETIKE)"
    wb.properties.title = path.stem
    wb.properties.created = FIXED_DT
    wb.properties.lastModifiedBy = "Sportel"
    buf = io.BytesIO()
    wb.save(buf)
    stamp = FIXED_DT.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(buf.getvalue())) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>" + stamp + rb"\g<2>",
                    data,
                )
            zi = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            dst.writestr(zi, data)
    path.write_bytes(out.getvalue())


def write_csv(path: Path, rows: list[list[str]], *, delimiter: str, bom: bool) -> None:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=delimiter, lineterminator="\r\n")
    writer.writerows(rows)
    path.write_bytes(buf.getvalue().encode("utf-8-sig" if bom else "utf-8"))


# --------------------------------------------------------------------------------------------
# 01 — citizen requests
# --------------------------------------------------------------------------------------------

CATEGORIES = [
    # category, weight, sla days, department
    ("Ndriçimi publik", 12, 10, "Drejtoria e Punëve Publike"),
    ("Kanalizimet", 10, 7, "Drejtoria e Punëve Publike"),
    ("Rrugë dhe trotuare", 14, 20, "Drejtoria e Punëve Publike"),
    ("Pastrimi dhe mbetjet", 16, 5, "Drejtoria e Shërbimeve Publike"),
    ("Gjelbërimi", 6, 10, "Drejtoria e Shërbimeve Publike"),
    ("Ujësjellësi", 11, 5, "Drejtoria e Shërbimeve Publike"),
    ("Leje dhe dokumente", 9, 30, "Drejtoria e Planifikimit dhe Zhvillimit të Territorit"),
    ("Ndihmë sociale", 8, 15, "Drejtoria e Shërbimeve Sociale"),
    ("Taksa dhe tarifa", 9, 10, "Drejtoria e të Ardhurave Vendore"),
    ("Transporti publik", 5, 15, "Drejtoria e Transportit dhe Lëvizshmërisë"),
]
PUBLIC_WORKS = "Drejtoria e Punëve Publike"
CHANNELS = [("Sportel", 38), ("Online", 27), ("Telefon", 25), ("E-mail", 10)]
REQUESTS_PER_MONTH = [262, 248, 296, 301, 318, 309, 329, 337]  # = 2,400


def gen_requests(rng: random.Random) -> list[dict]:
    unit_names = [u[0] for u in ADMIN_UNITS]
    rural_total = sum(u[1] for u in ADMIN_UNITS[1:])
    unit_weights = [58.0] + [42.0 * u[1] / rural_total for u in ADMIN_UNITS[1:]]
    rows: list[dict] = []
    seq = 0
    for month, n in zip(MONTHS, REQUESTS_PER_MONTH, strict=True):
        days_in_month = ((month.replace(day=28) + dt.timedelta(days=4)).replace(day=1) - month).days
        for _ in range(n):
            seq += 1
            while True:
                created = month + dt.timedelta(days=rng.randrange(days_in_month))
                if created.weekday() != 6 or rng.random() < 0.15:  # few Sunday entries
                    break
            if created > AS_OF:
                created = AS_OF
            cat, _, sla, dept = pick_weighted(rng, CATEGORIES, [c[1] for c in CATEGORIES])
            unit = pick_weighted(rng, unit_names, unit_weights)
            channel = pick_weighted(rng, [c[0] for c in CHANNELS], [c[1] for c in CHANNELS])

            summer_backlog = dept == PUBLIC_WORKS and created.month in (7, 8)
            on_time_p = 0.45 if summer_backlog else 0.9
            refused = rng.random() < 0.03
            stuck = created.month <= 6 and rng.random() < 0.018
            if refused:
                duration = rng.randint(1, max(2, sla))
            elif rng.random() < on_time_p:
                duration = rng.randint(max(1, sla // 5), sla)
            else:
                duration = sla + rng.randint(1, max(3, sla))
            closed = created + dt.timedelta(days=duration)
            if stuck or closed > AS_OF:
                closed = None
            if closed is None:
                age = (AS_OF - created).days
                status = "E re" if age <= 7 and rng.random() < 0.6 else "Në proces"
            else:
                status = "E refuzuar" if refused else "E mbyllur"
            rows.append(
                {
                    "nr": f"KQ-2026-{seq:05d}",
                    "created": created,
                    "name": fake_name(rng),
                    "phone": fake_phone(rng),
                    "category": cat,
                    "department": dept,
                    "unit": unit_spelling(rng, unit),
                    "unit_canonical": unit,
                    "channel": channel,
                    "status": status,
                    "closed": closed,
                    "sla": sla,
                }
            )
    return rows


def write_requests(path: Path, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kërkesat"
    headers = [
        "Nr.",
        "Data e regjistrimit",
        "Emri i kërkuesit",
        "Nr. telefoni",
        "Kategoria",
        "Drejtoria përgjegjëse",
        "Njësia administrative",
        "Kanali",
        "Statusi",
        "Data e mbylljes",
        "Afati (ditë)",
    ]
    write_titles(
        ws,
        [
            "TË DHËNA SINTETIKE (SINTETIKE) — Regjistri i kërkesave qytetare, janar–gusht 2026",
            "Sektori i Marrëdhënieve me Qytetarët (rol i supozuar) · Eksporti më 31.08.2026 · "
            + SYNTH_NOTE,
        ],
        len(headers),
    )
    write_header(ws, 3, headers)
    r = 4
    for row in rows:
        values = [
            row["nr"],
            fmt_date(row["created"]),  # text dates, as many exports store them
            row["name"],
            row["phone"],
            row["category"],
            row["department"],
            row["unit"],
            row["channel"],
            row["status"],
            row["closed"],  # real Excel dates
            row["sla"],
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if c == 10 and v is not None:
                cell.number_format = "DD.MM.YYYY"
        r += 1
    r += 1  # blank row before the total
    ws.cell(row=r, column=1, value="TOTALI")
    ws.cell(row=r, column=2, value=f"{len(rows)} kërkesa")
    style_total(ws, r, len(headers))
    set_widths(ws, [15, 14, 16, 18, 22, 44, 20, 10, 12, 14, 10])
    ws.freeze_panes = "A4"
    save_xlsx(wb, path)


# --------------------------------------------------------------------------------------------
# 02 — budget execution (thousand lek)
# --------------------------------------------------------------------------------------------

PROGRAMMES = [
    # code, name, annual current plan, annual capital plan (000 lek)
    ("01110", "Planifikim, menaxhim dhe administrim", 900_000, 60_000),
    ("03140", "Mbrojtja civile", 45_000, 15_000),
    ("04520", "Rrugët rurale", 40_000, 420_000),
    ("05100", "Menaxhimi i mbetjeve", 360_000, 40_000),
    ("06200", "Planifikimi urban", 70_000, 180_000),
    ("06440", "Ndriçimi publik", 110_000, 90_000),
    ("08130", "Sporti dhe argëtimi", 60_000, 70_000),
    ("09120", "Arsimi bazë", 820_000, 350_000),
    ("10400", "Kujdesi social për familjet dhe fëmijët", 300_000, 30_000),
]
CAPITAL_PLAN_PHASING = [0.03, 0.04, 0.06, 0.07, 0.08, 0.09, 0.10, 0.10]
CAPITAL_ACTUAL_SHAPE = [0.10, 0.18, 0.28, 0.35, 0.40, 0.42, 0.45, 0.44]
CAPITAL_EXECUTION_TARGET = 0.37
WASTE_CODE = "05100"
WASTE_SPIKE_MONTH = 5  # May: annual contract advance paid in one month


def gen_budget(rng: random.Random) -> list[dict]:
    rows: list[dict] = []
    for mi, month in enumerate(MONTHS):
        for code, name, cur_annual, cap_annual in PROGRAMMES:
            cur_plan = cur_annual / 12 * rng.uniform(0.97, 1.03)
            cur_actual = cur_plan * rng.uniform(0.88, 0.99)
            if code == WASTE_CODE:
                cur_actual = cur_plan * rng.uniform(0.94, 0.99)
                if month.month == WASTE_SPIKE_MONTH:
                    cur_actual = cur_plan * 1.95
            cap_plan = cap_annual * CAPITAL_PLAN_PHASING[mi] * rng.uniform(0.9, 1.1)
            cap_actual = cap_plan * CAPITAL_ACTUAL_SHAPE[mi] * rng.uniform(0.8, 1.2)
            rows.append(
                dict(
                    month=month,
                    code=code,
                    name=name,
                    kind="Korrente",
                    plan=cur_plan,
                    actual=cur_actual,
                )
            )
            rows.append(
                dict(
                    month=month,
                    code=code,
                    name=name,
                    kind="Kapitale",
                    plan=cap_plan,
                    actual=cap_actual,
                )
            )
    # calibrate capital execution to the planted ratio, then round to 0.1 thousand lek
    cap = [r for r in rows if r["kind"] == "Kapitale"]
    ratio = sum(r["actual"] for r in cap) / sum(r["plan"] for r in cap)
    for r in cap:
        r["actual"] *= CAPITAL_EXECUTION_TARGET / ratio
    for r in rows:
        r["plan_t"], r["actual_t"] = tenths(r["plan"]), tenths(r["actual"])
    return rows


def write_budget(path: Path, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Zbatimi"
    headers = [
        "Muaji",
        "Kodi i programit",
        "Programi",
        "Lloji i shpenzimit",
        "Plani (000 lekë)",
        "Fakti (000 lekë)",
    ]
    write_titles(
        ws,
        [
            "TË DHËNA SINTETIKE (SINTETIKE) — Zbatimi i buxhetit sipas programeve, "
            "janar–gusht 2026",
            "Drejtoria e Financës dhe Buxhetit (rol i supozuar) · Vlerat në mijë lekë · "
            "Gjendja më 31.08.2026 · " + SYNTH_NOTE,
        ],
        len(headers),
    )
    write_header(ws, 3, headers)
    r = 4
    for row in rows:
        values = [
            f"{MONTH_SQ[row['month'].month - 1]} {row['month'].year}",
            row["code"],  # text: keeps the leading zero
            row["name"],
            row["kind"],
            row["plan_t"] / 10,
            row["actual_t"] / 10,
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if c >= 5:
                cell.number_format = "#,##0.0"
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="TOTALI")
    ws.cell(row=r, column=5, value=sum(x["plan_t"] for x in rows) / 10).number_format = "#,##0.0"
    ws.cell(row=r, column=6, value=sum(x["actual_t"] for x in rows) / 10).number_format = "#,##0.0"
    style_total(ws, r, len(headers))
    set_widths(ws, [14, 16, 42, 18, 18, 18])
    ws.freeze_panes = "A4"
    save_xlsx(wb, path)


# --------------------------------------------------------------------------------------------
# Population reference (synthetic)
# --------------------------------------------------------------------------------------------


def write_population(path: Path) -> None:
    rows = [
        [
            "TË DHËNA SINTETIKE (SINTETIKE) — Popullsia sipas njësive administrative · "
            + SYNTH_NOTE,
            "",
            "",
        ],
        ["Njësia administrative", "Baza", "Banorë"],
    ]
    for name, census, _ in ADMIN_UNITS:
        rows.append([name, "census_2023", str(census)])
    for name, _, registry in ADMIN_UNITS:
        rows.append([name, "civil_registry", str(registry)])
    write_csv(path, rows, delimiter=",", bom=False)


# --------------------------------------------------------------------------------------------
# Envelope 1 — waste collection (CSV, ';', BOM, decimal comma)
# --------------------------------------------------------------------------------------------

SEASON = [0.95, 0.93, 0.98, 1.0, 1.02, 1.05, 1.09, 1.1]


def gen_waste(rng: random.Random) -> list[dict]:
    rows: list[dict] = []
    for mi, month in enumerate(MONTHS):
        days = ((month.replace(day=28) + dt.timedelta(days=4)).replace(day=1) - month).days
        for name, census, _ in ADMIN_UNITS:
            city = name == "Elbasan"
            kg_day = 1.05 if city else 0.55
            tonnes = census * kg_day * days / 1000 * SEASON[mi] * rng.uniform(0.96, 1.04)
            trips = round(tonnes / (6.5 if city else 3.4) * rng.uniform(0.95, 1.05))
            households = round(census / 3.4 * (0.97 if city else 0.72) * rng.uniform(0.99, 1.01))
            if name == "Gjinar" and month.month == 7:
                tonnes = 0.0  # planted placeholder: quantity missing, trips/households reported
            rows.append(
                dict(
                    month=month,
                    unit=unit_spelling(rng, name),
                    unit_canonical=name,
                    tonnes_t=tenths(tonnes),
                    trips=max(trips, 1),
                    households=households,
                )
            )
    return rows


def waste_csv_rows(rows: list[dict], *, drift: bool) -> list[list[str]]:
    ncols = 6 if drift else 5
    title = "TË DHËNA SINTETIKE (SINTETIKE) — Pastrimi dhe grumbullimi i mbetjeve, janar–gusht 2026"
    if drift:
        title += " (versioni 2)"
    out: list[list[str]] = [[title + " · " + SYNTH_NOTE] + [""] * (ncols - 1)]
    header = ["Muaji", "Njësia administrative", "Tonazhi" if drift else "Sasia (ton)",
              "Nr. i kursimeve", "Familje të mbuluara"]  # fmt: skip
    if drift:
        header.append("Operatori")
    out.append(header)
    for r in rows:
        line = [
            r["month"].strftime("%Y-%m"),
            r["unit"],
            fmt_sq(r["tonnes_t"] / 10, 1) if r["tonnes_t"] else "0",
            fmt_sq(r["trips"]),
            fmt_sq(r["households"]),
        ]
        if drift:
            line.append("Operatori A" if r["unit_canonical"] == "Elbasan" else "Operatori B")
        out.append(line)
    out.append([""] * ncols)
    total = ["TOTALI", "", fmt_sq(sum(r["tonnes_t"] for r in rows) / 10, 1),
             fmt_sq(sum(r["trips"] for r in rows)), ""]  # fmt: skip
    if drift:
        total.append("")
    out.append(total)
    return out


# --------------------------------------------------------------------------------------------
# Envelope 2 — local revenue collection (thousand lek)
# --------------------------------------------------------------------------------------------

REVENUE_TYPES = [
    # type, payer, monthly plan (000 lek), collection ratio
    ("Taksa e ndërtesës", "Familje", 5_000, 0.70),
    ("Taksa e ndërtesës", "Biznes", 11_500, 0.88),
    ("Tarifa e pastrimit", "Familje", 12_500, None),
    ("Tarifa e pastrimit", "Biznes", 15_800, None),
    ("Taksa e ndikimit në infrastrukturë", "Familje", 3_300, 0.95),
    ("Taksa e ndikimit në infrastrukturë", "Biznes", 15_000, 0.92),
    ("Taksa e truallit", "Familje", 1_700, 0.60),
    ("Taksa e truallit", "Biznes", 2_900, 0.80),
    ("Gjoba", "Familje", 700, 0.55),
    ("Gjoba", "Biznes", 1_800, 0.70),
    ("Qira", "Familje", 400, 0.90),
    ("Qira", "Biznes", 3_750, 0.85),
]
FEE_COVERAGE_TARGET = 0.55


def gen_revenue(rng: random.Random, budget_rows: list[dict]) -> list[dict]:
    waste_spend = {m: 0 for m in MONTHS}
    for b in budget_rows:
        if b["code"] == WASTE_CODE:
            waste_spend[b["month"]] += b["actual_t"]
    total_waste = sum(waste_spend.values()) / 10  # 000 lek
    fee_target = total_waste * FEE_COVERAGE_TARGET
    rows: list[dict] = []
    for month in MONTHS:
        for rtype, payer, plan, ratio in REVENUE_TYPES:
            p = plan * rng.uniform(0.95, 1.05)
            if ratio is None:
                share = 0.45 if payer == "Familje" else 0.55
                collected = fee_target / len(MONTHS) * share * rng.uniform(0.85, 1.15)
            else:
                collected = p * ratio * rng.uniform(0.9, 1.1)
            rows.append(dict(month=month, type=rtype, payer=payer, plan=p, collected=collected))
    fee = [r for r in rows if r["type"] == "Tarifa e pastrimit"]
    scale = fee_target / sum(r["collected"] for r in fee)
    for r in fee:
        r["collected"] *= scale
    for r in rows:
        r["plan_t"], r["collected_t"] = tenths(r["plan"]), tenths(r["collected"])
    return rows


def write_revenue(path: Path, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Arkëtimi"
    headers = [
        "Muaji",
        "Lloji i të ardhurës",
        "Kategoria e paguesit",
        "Plani (000 lekë)",
        "Arkëtuar (000 lekë)",
    ]
    ws.cell(row=1, column=1, value=(
        "TË DHËNA SINTETIKE (SINTETIKE) — Taksat dhe tarifat vendore: plani dhe arkëtimi, 2026 · "
        + SYNTH_NOTE
    ))  # fmt: skip
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(row=1, column=1).font = Font(bold=True, size=12, color="7F6000")
    ws.cell(row=1, column=1).fill = TITLE_FILL
    # merged header band over the money columns
    ws.cell(row=2, column=4, value="Vlera në mijë lekë (000 lekë)")
    ws.merge_cells(start_row=2, start_column=4, end_row=2, end_column=5)
    band = ws.cell(row=2, column=4)
    band.font = Font(bold=True)
    band.fill = HEADER_FILL
    band.alignment = Alignment(horizontal="center")
    write_header(ws, 3, headers)
    r = 4
    for row in rows:
        ws.cell(row=r, column=1, value=row["month"]).number_format = "MM.YYYY"
        ws.cell(row=r, column=2, value=row["type"])
        ws.cell(row=r, column=3, value=row["payer"])
        ws.cell(row=r, column=4, value=row["plan_t"] / 10).number_format = "#,##0.0"
        ws.cell(row=r, column=5, value=row["collected_t"] / 10).number_format = "#,##0.0"
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Gjithsej")
    ws.cell(row=r, column=4, value=sum(x["plan_t"] for x in rows) / 10).number_format = "#,##0.0"
    ws.cell(
        row=r, column=5, value=sum(x["collected_t"] for x in rows) / 10
    ).number_format = "#,##0.0"
    style_total(ws, r, len(headers))
    set_widths(ws, [12, 36, 20, 18, 20])
    ws.freeze_panes = "A4"
    save_xlsx(wb, path)


# --------------------------------------------------------------------------------------------
# Envelope 3 — human resources (headcount at 31.08.2026)
# --------------------------------------------------------------------------------------------

DIRECTORATES = [
    # synthetic structure; headcount, hires 2026, leavers 2026
    ("Kabineti i Kryetarit dhe Sekretaria", 40, 3, 4),
    ("Drejtoria e Financës dhe Buxhetit", 38, 2, 3),
    ("Drejtoria e të Ardhurave Vendore", 52, 5, 4),
    ("Drejtoria e Burimeve Njerëzore", 14, 1, 1),
    ("Drejtoria e Shërbimeve Publike", 380, 31, 27),
    ("Drejtoria e Punëve Publike", 145, 9, 12),
    ("Drejtoria e Planifikimit dhe Zhvillimit të Territorit", 41, 3, 5),
    ("Drejtoria e Shërbimeve Sociale", 96, 6, 7),
    ("Drejtoria e Arsimit, Kulturës dhe Sportit", 330, 18, 14),
    ("Drejtoria e Transportit dhe Lëvizshmërisë", 28, 2, 2),
    ("Drejtoria e Mbrojtjes Civile dhe Emergjencave", 72, 4, 3),
    ("Policia Bashkiake", 85, 7, 6),
    ("Njësitë administrative (12)", 180, 8, 10),
]


def write_staff(path: Path, rng: random.Random) -> list[dict]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Personeli"
    headers = [
        "Nr.",
        "Drejtoria",
        "Numri i punonjësve (31.08.2026)",
        "Pranime 2026",
        "Largime 2026",
        "Përgjegjësi i drejtorisë",
    ]
    write_titles(
        ws,
        [
            "TË DHËNA SINTETIKE (SINTETIKE) — Burimet njerëzore sipas drejtorive, 2026",
            "Drejtoria e Burimeve Njerëzore (rol i supozuar) · Gjendja më 31.08.2026 · "
            + SYNTH_NOTE,
        ],
        len(headers),
    )
    write_header(ws, 3, headers)
    rows = []
    r = 4
    for i, (name, hc, hires, leavers) in enumerate(DIRECTORATES, start=1):
        values = [i, name, hc, hires, leavers, fake_name(rng)]
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v)
        rows.append(dict(name=name, headcount=hc, hires=hires, leavers=leavers))
        r += 1
    r += 1
    ws.cell(row=r, column=2, value="Gjithsej")
    ws.cell(row=r, column=3, value=sum(x["headcount"] for x in rows))
    ws.cell(row=r, column=4, value=sum(x["hires"] for x in rows))
    ws.cell(row=r, column=5, value=sum(x["leavers"] for x in rows))
    style_total(ws, r, len(headers))
    set_widths(ws, [6, 50, 18, 14, 14, 24])
    ws.freeze_panes = "A4"
    save_xlsx(wb, path)
    return rows


# --------------------------------------------------------------------------------------------
# Manifest and planted-pattern checks
# --------------------------------------------------------------------------------------------

FILES = [
    # name, label sq, label en, dataset_hint, envelope, preload
    ("01_SINTETIKE_kerkesat_qytetare_jan-gus_2026.xlsx", "Kërkesat qytetare, jan–gus 2026",
     "Citizen requests, Jan–Aug 2026", "requests", None, True),
    ("02_SINTETIKE_zbatimi_buxhetit_jan-gus_2026.xlsx", "Zbatimi i buxhetit, jan–gus 2026",
     "Budget execution, Jan–Aug 2026", "budget", None, True),
    ("ref_SINTETIKE_popullsia_njesite.csv", "Popullsia sipas njësive (referencë)",
     "Population by administrative unit (reference)", "population", None, True),
    ("zarfi-1_SINTETIKE_pastrimi_mbetjet_2026.csv", "Zarfi 1 · Pastrimi dhe mbetjet 2026",
     "Envelope 1 · Cleaning and waste 2026", "waste", 1, False),
    ("zarfi-2_SINTETIKE_taksat_tarifat_arketimi_2026.xlsx",
     "Zarfi 2 · Taksat dhe tarifat, arkëtimi 2026",
     "Envelope 2 · Taxes and fees, collection 2026", "revenue", 2, False),
    ("zarfi-3_SINTETIKE_burimet_njerezore_2026.xlsx", "Zarfi 3 · Burimet njerëzore 2026",
     "Envelope 3 · Human resources 2026", "staff", 3, False),
    ("drift_SINTETIKE_pastrimi_mbetjet_2026_v2.csv",
     "Pastrimi dhe mbetjet v2 (format i ndryshuar)",
     "Cleaning and waste v2 (changed format)", "waste", None, False),
]  # fmt: skip


def write_manifest(path: Path) -> None:
    manifest = {
        "generated_by": "api/scripts/generate_samples.py",
        "seed": SEED,
        "as_of": AS_OF.isoformat(),
        "synthetic": True,
        "note": {
            "sq": "Eksporte sintetike departamentesh për demonstrim; jo të dhëna zyrtare.",
            "en": "Synthetic department exports for demonstration; not official data.",
        },
        "files": [
            {
                "name": name,
                "label": {"sq": sq, "en": en},
                "dataset_hint": hint,
                "envelope": envelope,
                "preload": preload,
                "synthetic": True,
            }
            for name, sq, en, hint, envelope, preload in FILES
        ],
    }
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def check_patterns(requests, budget, waste, revenue, staff) -> dict:
    """Assert the planted, discoverable patterns of contract §1 and return a summary."""
    summary: dict = {}
    assert len(requests) == 2_400
    assert sum(u[1] for u in ADMIN_UNITS) == CENSUS_TOTAL
    assert sum(u[2] for u in ADMIN_UNITS) == REGISTRY_TOTAL

    # public works on-time rate by closing month drops in July-August
    on_time = defaultdict(lambda: [0, 0])
    for r in requests:
        if r["closed"] and r["department"] == PUBLIC_WORKS:
            k = on_time[r["closed"].month]
            k[1] += 1
            k[0] += (r["closed"] - r["created"]).days <= r["sla"]
    pw = {m: round(100 * a / b, 1) for m, (a, b) in sorted(on_time.items())}
    summary["public_works_on_time_pct_by_month"] = pw
    assert min(pw[m] for m in range(1, 7)) > 80 and max(pw[7], pw[8]) < 72, pw

    cap = [b for b in budget if b["kind"] == "Kapitale"]
    cap_pct = 100 * sum(b["actual_t"] for b in cap) / sum(b["plan_t"] for b in cap)
    summary["capital_execution_pct"] = round(cap_pct, 1)
    assert 35 <= cap_pct <= 40, cap_pct
    total_pct = 100 * sum(b["actual_t"] for b in budget) / sum(b["plan_t"] for b in budget)
    summary["budget_execution_pct"] = round(total_pct, 1)

    gj = [w for w in waste if w["unit_canonical"] == "Gjinar" and w["month"].month == 7]
    assert gj and gj[0]["tonnes_t"] == 0

    tonnes = defaultdict(int)
    for w in waste:
        tonnes[w["month"]] += w["tonnes_t"]
    spend = defaultdict(int)
    for b in budget:
        if b["code"] == WASTE_CODE:
            spend[b["month"]] += b["actual_t"]
    cpt = [spend[m] * 1000 / tonnes[m] for m in MONTHS]  # lek per tonne
    swings = [abs(cpt[i] / cpt[i - 1] - 1) for i in range(1, len(cpt))]
    summary["cost_per_tonne_lek"] = [round(x) for x in cpt]
    assert max(swings) > 0.5, swings

    fee = sum(r["collected_t"] for r in revenue if r["type"] == "Tarifa e pastrimit")
    coverage = 100 * fee / sum(spend.values())
    summary["cleaning_fee_coverage_pct"] = round(coverage, 1)
    assert 50 <= coverage <= 60, coverage
    summary["revenue_collection_pct"] = round(
        100 * sum(r["collected_t"] for r in revenue) / sum(r["plan_t"] for r in revenue), 1
    )
    summary["waste_tonnes_ytd"] = sum(tonnes.values()) / 10
    summary["headcount"] = sum(s["headcount"] for s in staff)
    assert 1_450 <= summary["headcount"] <= 1_550
    return summary


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)

    requests = gen_requests(random.Random(SEED + 1))
    write_requests(OUT / FILES[0][0], requests)

    budget = gen_budget(random.Random(SEED + 2))
    write_budget(OUT / FILES[1][0], budget)

    write_population(OUT / FILES[2][0])

    waste = gen_waste(random.Random(SEED + 3))
    write_csv(OUT / FILES[3][0], waste_csv_rows(waste, drift=False), delimiter=";", bom=True)
    write_csv(OUT / FILES[6][0], waste_csv_rows(waste, drift=True), delimiter=";", bom=True)

    revenue = gen_revenue(random.Random(SEED + 4), budget)
    write_revenue(OUT / FILES[4][0], revenue)

    staff = write_staff(OUT / FILES[5][0], random.Random(SEED + 5))

    write_manifest(OUT / "manifest.json")
    summary = check_patterns(requests, budget, waste, revenue, staff)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    for name, *_ in FILES:
        print(f"wrote {name} ({(OUT / name).stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
