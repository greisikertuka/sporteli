"""Independent reading of the raw synthetic samples, for the golden set.

Nothing here uses the ingest pipeline, the catalog parsers or the passports' SQL. Each file is
read with openpyxl/csv and file-specific code, and every expected indicator value is computed
in plain Python from the written formula definitions. The copilot's answers (passport SQL over
whatever the ingest pipeline loaded) are then compared with these numbers.

``read_samples()`` also returns canonical rows so tests can load the warehouse directly without
the ingest pipeline (``load_into``).
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

SAMPLES = Path(__file__).resolve().parents[2] / "samples"

REQUESTS_FILE = "01_SINTETIKE_kerkesat_qytetare_jan-gus_2026.xlsx"
BUDGET_FILE = "02_SINTETIKE_zbatimi_buxhetit_jan-gus_2026.xlsx"
POPULATION_FILE = "ref_SINTETIKE_popullsia_njesite.csv"
WASTE_FILE = "zarfi-1_SINTETIKE_pastrimi_mbetjet_2026.csv"
REVENUE_FILE = "zarfi-2_SINTETIKE_taksat_tarifat_arketimi_2026.xlsx"
STAFF_FILE = "zarfi-3_SINTETIKE_burimet_njerezore_2026.xlsx"

PRELOAD = ("requests", "budget", "population")
ENVELOPES = ("waste", "revenue", "staff")

_MONTHS_SQ = {
    "janar": 1,
    "shkurt": 2,
    "mars": 3,
    "prill": 4,
    "maj": 5,
    "qershor": 6,
    "korrik": 7,
    "gusht": 8,
    "shtator": 9,
    "tetor": 10,
    "nentor": 11,
    "dhjetor": 12,
}


def _plain(text: object) -> str:
    s = unicodedata.normalize("NFKD", str(text or ""))
    return "".join(c for c in s if not unicodedata.combining(c)).casefold().strip()


def _albanian_number(text: str) -> float:
    """ "2.562,4" → 2562.4; "23.944" → 23944 (thousands dot, decimal comma)."""
    return float(text.strip().replace(".", "").replace(",", "."))


def _xlsx(name: str, samples: Path) -> list[tuple]:
    ws = load_workbook(samples / name, data_only=True, read_only=True).active
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


def _csv(name: str, samples: Path, delimiter: str) -> list[list[str]]:
    text = (samples / name).read_bytes().decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _last_day(d: dt.date) -> dt.date:
    nxt = (d.replace(day=28) + dt.timedelta(days=4)).replace(day=1)
    return nxt - dt.timedelta(days=1)


@dataclass
class Samples:
    """Canonical rows per table (tuples in warehouse column order, without source_id)."""

    rows: dict[str, list[tuple]] = field(default_factory=dict)
    filenames: dict[str, str] = field(default_factory=dict)


def read_samples(samples: Path = SAMPLES) -> Samples:
    out = Samples()

    # requests: header on row 3; data rows start with "KQ-"
    rows = _xlsx(REQUESTS_FILE, samples)
    header = [str(h or "").strip() for h in rows[2]]
    col = {h: i for i, h in enumerate(header)}
    req = []
    for row_no, r in enumerate(rows[3:], start=4):
        if not r[0] or not str(r[0]).startswith("KQ-"):
            continue
        created = dt.datetime.strptime(str(r[col["Data e regjistrimit"]]), "%d.%m.%Y").date()
        closed_raw = r[col["Data e mbylljes"]]
        closed = closed_raw.date() if isinstance(closed_raw, dt.datetime) else closed_raw
        sla = r[col["Afati (ditë)"]]
        req.append(
            (
                row_no,
                str(r[0]),
                created,
                closed,
                r[col["Kategoria"]],
                r[col["Drejtoria përgjegjëse"]],
                r[col["Njësia administrative"]],
                r[col["Kanali"]],
                r[col["Statusi"]],
                int(sla) if sla is not None else None,
            )
        )
    out.rows["request"], out.filenames["request"] = req, REQUESTS_FILE

    # budget: "Janar 2026" months, values in thousands of lek, TOTALI excluded
    rows = _xlsx(BUDGET_FILE, samples)
    bud = []
    for row_no, r in enumerate(rows[3:], start=4):
        if r[0] is None or _plain(r[0]) == "totali":
            continue
        name, year = _plain(r[0]).split()
        month = dt.date(int(year), _MONTHS_SQ[name], 1)
        kind = "capital" if _plain(r[3]).startswith("kapital") else "current"
        bud.append((row_no, month, str(r[1]), r[2], kind, r[4] * 1000.0, r[5] * 1000.0))
    out.rows["budget_line"], out.filenames["budget_line"] = bud, BUDGET_FILE

    # population: comma CSV, title row 1, header row 2
    rows = _csv(POPULATION_FILE, samples, ",")
    pop = [(i, r[0], r[1], int(r[2])) for i, r in enumerate(rows[2:], start=3) if r and r[0]]
    out.rows["population"], out.filenames["population"] = pop, POPULATION_FILE

    # waste: ';' CSV, decimal comma, TOTALI excluded
    rows = _csv(WASTE_FILE, samples, ";")
    wst = []
    for row_no, r in enumerate(rows[2:], start=3):
        if not r or not r[0] or _plain(r[0]) == "totali":
            continue
        year, month = (int(x) for x in r[0].split("-"))
        wst.append(
            (
                row_no,
                dt.date(year, month, 1),
                r[1],
                _albanian_number(r[2]),
                int(_albanian_number(r[3])),
                int(_albanian_number(r[4])),
            )
        )
    out.rows["waste_collection"], out.filenames["waste_collection"] = wst, WASTE_FILE

    # revenue: month cells are dates, thousands of lek, "Gjithsej" excluded
    rows = _xlsx(REVENUE_FILE, samples)
    rev = []
    for row_no, r in enumerate(rows[3:], start=4):
        if r[0] is None or _plain(r[0]) == "gjithsej":
            continue
        month = r[0].date().replace(day=1) if isinstance(r[0], dt.datetime) else r[0]
        rev.append((row_no, month, r[1], r[2], r[3] * 1000.0, r[4] * 1000.0))
    out.rows["revenue"], out.filenames["revenue"] = rev, REVENUE_FILE

    # staff: directorate in column B, snapshot date inside the headcount header
    rows = _xlsx(STAFF_FILE, samples)
    m = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", str(rows[2][2]))
    as_of = dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else None
    hr = []
    for row_no, r in enumerate(rows[3:], start=4):
        if r[1] is None or _plain(r[1]) == "gjithsej":
            continue
        hr.append((row_no, r[1], int(r[2]), int(r[3]), int(r[4]), as_of))
    out.rows["staff"], out.filenames["staff"] = hr, STAFF_FILE
    return out


TABLE_OF = {
    "requests": "request",
    "budget": "budget_line",
    "population": "population",
    "waste": "waste_collection",
    "revenue": "revenue",
    "staff": "staff",
}


def load_into(con, samples: Samples, datasets) -> None:
    """Insert the given datasets' rows (plus their ``source`` rows) straight into the warehouse.

    Admin units are normalised with the catalog so breakdowns group like the real pipeline.
    """
    from app.catalog import normalize_admin_unit
    from app.warehouse.db import insert_rows

    for ds in datasets:
        table = TABLE_OF[ds]
        sid = f"golden-{ds}"
        rows = []
        for r in samples.rows[table]:
            r = list(r)
            if table == "request":
                r[6] = normalize_admin_unit(r[6])
            elif table in ("waste_collection", "population"):
                idx = 2 if table == "waste_collection" else 1
                r[idx] = normalize_admin_unit(r[idx])
            rows.append((sid, *r))
        insert_rows(con, table, rows)
        con.execute(
            "INSERT INTO source (id, filename, file_hash, dataset, synthetic) "
            "VALUES (?, ?, ?, ?, true)",
            [sid, samples.filenames[table], f"hash-{ds}", ds],
        )


# --------------------------------------------------------------------------------------------
# Expected values, computed in plain Python from the written definitions
# --------------------------------------------------------------------------------------------


def expected_values(samples: Samples, basis: str = "census_2023") -> dict[str, float]:
    v: dict[str, float] = {}
    req = samples.rows["request"]
    # columns: row_no, id, created, closed, category, department, unit, channel, status, sla
    last_created = max(r[2] for r in req)
    y0, m1 = dt.date(last_created.year, 1, 1), _last_day(last_created)

    v["requests_received"] = float(sum(1 for r in req if y0 <= r[2] <= m1))
    closed = [r for r in req if r[3] is not None and y0 <= r[3] <= m1 and r[9] is not None]
    on_time = sum(1 for r in closed if (r[3] - r[2]).days <= r[9])
    v["requests_on_time_pct"] = 100.0 * on_time / len(closed)
    overdue = [
        r
        for r in req
        if r[2] <= m1
        and (r[3] is None or r[3] > m1)
        and r[9] is not None
        and (m1 - r[2]).days > r[9]
        and "refuz" not in _plain(r[8])
    ]
    v["requests_open_overdue"] = float(len(overdue))
    durations = [
        (r[3] - r[2]).days for r in req if r[3] is not None and y0 <= r[3] <= m1 and r[3] >= r[2]
    ]
    v["requests_avg_days"] = sum(durations) / len(durations)

    bud = samples.rows["budget_line"]
    # columns: row_no, month, code, programme, line_type, planned, actual
    b_last = max(r[1] for r in bud)
    b_rows = [r for r in bud if dt.date(b_last.year, 1, 1) <= r[1] <= b_last]
    v["budget_execution_pct"] = 100.0 * sum(r[6] for r in b_rows) / sum(r[5] for r in b_rows)
    cap = [r for r in b_rows if r[4] == "capital"]
    v["capital_execution_pct"] = 100.0 * sum(r[6] for r in cap) / sum(r[5] for r in cap)

    residents = sum(r[3] for r in samples.rows["population"] if r[2] == basis)

    wst = samples.rows.get("waste_collection") or []
    if wst:
        w_last = max(r[1] for r in wst)
        w_rows = [r for r in wst if dt.date(w_last.year, 1, 1) <= r[1] <= w_last]
        tonnes = sum(r[3] for r in w_rows)
        months = len({r[1] for r in w_rows})
        v["waste_tonnes"] = tonnes
        v["waste_kg_per_resident"] = tonnes * 1000.0 / residents * 12.0 / months
        common = min(w_last, b_last)
        c0 = dt.date(common.year, 1, 1)
        waste_cost = sum(r[6] for r in bud if r[2] == "05100" and c0 <= r[1] <= common)
        waste_t = sum(r[3] for r in wst if c0 <= r[1] <= common)
        v["waste_cost_per_tonne"] = waste_cost / waste_t

    rev = samples.rows.get("revenue") or []
    if rev:
        r_last = max(r[1] for r in rev)
        r_rows = [r for r in rev if dt.date(r_last.year, 1, 1) <= r[1] <= r_last]
        v["revenue_collection_pct"] = 100.0 * sum(r[5] for r in r_rows) / sum(r[4] for r in r_rows)
        common = min(r_last, b_last)
        c0 = dt.date(common.year, 1, 1)
        fee = sum(r[5] for r in rev if "pastrim" in _plain(r[2]) and c0 <= r[1] <= common)
        cost = sum(r[6] for r in bud if r[2] == "05100" and c0 <= r[1] <= common)
        v["cleaning_fee_coverage_pct"] = 100.0 * fee / cost

    hr = samples.rows.get("staff") or []
    if hr:
        latest = max(r[5] for r in hr)
        snap = [r for r in hr if r[5] == latest]
        hc = sum(r[2] for r in snap)
        hires = sum(r[3] for r in snap)
        leavers = sum(r[4] for r in snap)
        v["staff_per_1000"] = hc * 1000.0 / residents
        v["staff_turnover_pct"] = 100.0 * leavers / ((hc + (hc - hires + leavers)) / 2.0)
    return v
