"""Regression tests for review findings outside the ingest supersede/reconcile path: error
shape and CORS on unexpected errors, request size guard, sandbox result caps, USING SAMPLE,
PII in set-aside rows, per-column number style, UTF-16 CSVs, zip caps, export notes and the
additive indicator fields."""

import io
import zipfile
from pathlib import Path

import duckdb
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import catalog
from app.indicators import registry as reg
from app.ingest import pipeline, reader
from app.ingest.errors import IngestError
from app.sqlguard.guard import check_sql
from app.sqlguard.sandbox import MAX_CELL_CHARS, MAX_RESULT_BYTES, run_sandboxed

SAMPLES = Path(__file__).resolve().parents[1] / "samples"
WASTE = "zarfi-1_SINTETIKE_pastrimi_mbetjet_2026.csv"
ORIGIN = "http://localhost:3000"


# ---------------------------------------------------------------- errors, CORS, size guard


@pytest.fixture
def app_client(settings_env):
    from app.main import create_app

    app = create_app()

    @app.get("/boom")
    def boom() -> dict:
        raise RuntimeError("unexpected")

    @app.get("/boom-db")
    def boom_db() -> dict:
        raise duckdb.CatalogException("Table with name request does not exist!")

    with TestClient(app) as c:
        yield c


def test_unhandled_error_uses_the_contract_shape_with_cors(app_client):
    r = app_client.get("/boom", headers={"Origin": ORIGIN})
    assert r.status_code == 500
    assert r.headers["content-type"].startswith("application/json")
    assert r.headers.get("access-control-allow-origin") == ORIGIN
    assert r.json() == {
        "detail": {
            "code": "internal_error",
            "message": {"sq": "Ndodhi një gabim i brendshëm.", "en": "An internal error occurred."},
        }
    }


def test_database_error_uses_the_contract_shape_with_cors(app_client):
    r = app_client.get("/boom-db", headers={"Origin": ORIGIN})
    assert r.status_code == 500
    assert r.headers.get("access-control-allow-origin") == ORIGIN
    detail = r.json()["detail"]
    assert detail["code"] == "database_error"
    assert set(detail["message"]) == {"sq", "en"}


def test_oversized_request_is_refused_by_content_length(app_client):
    r = app_client.post(
        "/api/v1/ingest/preview",
        content=b"x",
        headers={
            "Origin": ORIGIN,
            "Content-Length": str(17 * 1024 * 1024),
            "Content-Type": "multipart/form-data; boundary=zzz",
        },
    )
    assert r.status_code == 413
    assert r.json()["detail"]["code"] == "file_too_large"
    assert r.headers.get("access-control-allow-origin") == ORIGIN


def test_oversized_streamed_body_is_refused(settings_env):
    from app.meta.errors import GuardMiddleware

    received: list[int] = []

    async def app(scope, receive, send):
        while True:
            msg = await receive()
            received.append(len(msg.get("body", b"")))
            if not msg.get("more_body"):
                break
        await send({"type": "http.response.start", "status": 200, "headers": []})
        await send({"type": "http.response.body", "body": b"ok"})

    guarded = GuardMiddleware(app, max_bytes=10)
    chunks = [b"a" * 8, b"b" * 8, b"c" * 8]
    sent: list[dict] = []

    async def receive():
        body = chunks.pop(0)
        return {"type": "http.request", "body": body, "more_body": bool(chunks)}

    async def send(message):
        sent.append(message)

    import asyncio

    asyncio.run(guarded({"type": "http", "headers": []}, receive, send))
    assert sent[0]["status"] == 413
    assert received == [8]  # stopped at the chunk that crossed the limit


# ---------------------------------------------------------------- sandbox and guard


def test_exploratory_results_are_size_capped(db):
    verdict = check_sql("SELECT repeat('x', 100000) AS a, repeat('y', 100000) AS b")
    assert verdict.allowed
    res = run_sandboxed(db, verdict.sql, verdict.tables)
    assert len(res.rows) == 1
    assert len(res.rows[0][0]) == MAX_CELL_CHARS and res.rows[0][0].endswith("…")

    pipeline.ingest_path(SAMPLES / WASTE, "waste", con=db)  # 104 rows
    cols = ", ".join(f"repeat('z', 20000) AS c{i}" for i in range(24))
    verdict = check_sql(f"SELECT {cols} FROM waste_collection")
    assert verdict.allowed
    res = run_sandboxed(db, verdict.sql, verdict.tables)
    assert res.truncated is True
    assert 0 < len(res.rows) < 104
    assert sum(len(c) for r in res.rows for c in r) <= MAX_RESULT_BYTES


def test_ask_with_a_huge_string_returns_a_small_response(client):
    r = client.post(
        "/api/v1/ask",
        json={"question": "SELECT repeat('x', 10000000) AS a, repeat('y', 10000000) AS b"},
    )
    assert r.status_code == 200
    assert len(r.content) < 100_000


def test_using_sample_is_wrapped_and_runs(db):
    pipeline.ingest_path(SAMPLES / WASTE, "waste", con=db)
    for sql, n in [
        ("SELECT * FROM waste_collection USING SAMPLE 10", 10),
        ("SELECT * FROM waste_collection USING SAMPLE 10 LIMIT 4", 4),
    ]:
        verdict = check_sql(sql)
        assert verdict.allowed, verdict
        assert verdict.sql.rstrip().endswith(f"LIMIT {200 if n == 10 else n}")
        res = run_sandboxed(db, verdict.sql, verdict.tables)
        assert len(res.rows) == n


# ---------------------------------------------------------------- PII in set-aside rows


def test_excluded_row_texts_carry_no_personal_values(db):
    text = (
        "Muaji;Njësia administrative;Sasia (ton);Kontakti\n"
        "2026-07;Gjinar;21,4;arben.kola@example.com\n"
        "2026-07;Gracen;26,0;mira.dema@example.com\n"
        "2026-07;Funarë;18,5;ilir.hoxha@example.com\n"
        "Totali për Arben Kola +355 69 123 45 67;;65,9;arben.kola@example.com\n"
        "Përgatiti: Arben Kola, tel. 069 123 4567, email arben.kola@example.com;;;\n"
    )
    p = pipeline.preview(text.encode("utf-8"), "kontakt_SINTETIKE.csv", con=db, use_llm=False)
    assert any(c["pii"] == "email" and c["dropped"] for c in p["columns"])
    blob = " ".join(e["text"] for e in p["excluded_rows"])
    assert "example.com" not in blob
    assert "123 45 67" not in blob and "123 4567" not in blob
    assert "[telefon]" in blob and "[email]" in blob
    assert "65,9" in blob or "65.9" in blob  # amounts stay
    steps = " ".join(s["message"]["en"] + s["message"]["sq"] for s in p["steps"])
    assert "example.com" not in steps and "123 45 67" not in steps


# ---------------------------------------------------------------- numbers, encodings, zips


def test_number_style_is_decided_per_column():
    assert catalog.number_style(["2.12", "2.125", "12.345"]) == "en"
    assert catalog.number_style(["1,234.5", "1,234"]) == "en"
    assert catalog.number_style(["4.683,3", "12.345"]) == "sq"
    assert catalog.number_style(["1.234", "5.678"]) is None
    assert catalog.parse_number("2.125", "en") == 2.125
    assert catalog.parse_number("12.345", "en") == 12.345
    assert catalog.parse_number("1,234", "en") == 1234
    assert catalog.parse_number("12.345", "sq") == 12345
    assert catalog.parse_number("1,5", "sq") == 1.5


def test_dot_decimal_column_is_read_at_one_scale(db):
    text = (
        "Muaji,Njësia administrative,Sasia (ton)\n"
        "2026-07,Gjinar,2.12\n"
        "2026-07,Gracen,2.125\n"
        "2026-07,Funarë,12.345\n"
    )
    receipt = pipeline.ingest_path(_tmp(db, "dot_SINTETIKE.csv", text), "waste", con=db)
    assert receipt["rows_loaded"] == 3
    got = dict(
        db.execute(
            "SELECT admin_unit, tonnes FROM waste_collection WHERE source_id = ?",
            [receipt["source_id"]],
        ).fetchall()
    )
    assert got == {"Gjinar": 2.12, "Gracen": 2.125, "Funarë": 12.345}


def _tmp(db, name: str, text: str) -> Path:
    path = Path(__file__).resolve().parent / "_tmp_review"
    path.mkdir(exist_ok=True)
    f = path / name
    f.write_text(text, encoding="utf-8")
    return f


@pytest.fixture(autouse=True)
def _cleanup_tmp():
    yield
    folder = Path(__file__).resolve().parent / "_tmp_review"
    if folder.exists():
        for f in folder.iterdir():
            f.unlink()
        folder.rmdir()


@pytest.mark.parametrize("encoding", ["utf-16", "utf-16-le", "utf-16-be"])
def test_utf16_csv_is_read_like_utf8(encoding):
    text = (SAMPLES / WASTE).read_bytes().decode("utf-8-sig")
    table = reader.read_table(text.encode(encoding), WASTE)
    assert table.encoding.startswith("utf-16")
    assert table.rows[1][:3] == ["Muaji", "Njësia administrative", "Sasia (ton)"]
    utf8 = reader.read_table(text.encode("utf-8"), WASTE)
    assert table.rows == utf8.rows


def test_xlsx_zip_cap_counts_decompressed_bytes(monkeypatch):
    data = (SAMPLES / "zarfi-3_SINTETIKE_burimet_njerezore_2026.xlsx").read_bytes()
    reader.read_table(data, "x.xlsx")  # fine at the real limit
    # an archive whose headers understate the size: rewrite one member's declared size
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/big.bin", b"\0" * 300_000)
    lying = bytearray(buf.getvalue())
    cd = lying.rfind(b"PK\x01\x02")
    lying[cd + 24 : cd + 28] = (10).to_bytes(4, "little")  # central directory: file_size = 10
    monkeypatch.setattr(reader, "MAX_UNZIPPED_BYTES", 100_000)
    with pytest.raises(IngestError):
        reader._check_zip(bytes(lying))


def test_xlsx_read_only_matches_merged_titles_and_rows():
    name = "02_SINTETIKE_zbatimi_buxhetit_jan-gus_2026.xlsx"
    table = reader.read_table((SAMPLES / name).read_bytes(), name)
    wb = load_workbook(SAMPLES / name, data_only=True)
    ws = wb.active
    assert len(table.rows) == ws.max_row
    assert sorted(table.merged) == sorted(
        (r.min_row - 1, r.min_col - 1, r.max_row - 1, r.max_col - 1) for r in ws.merged_cells.ranges
    )
    # a merged title is forward-filled across its range
    assert table.rows[0][0] == table.rows[0][5]


def test_xlsx_row_limit_stops_reading(monkeypatch):
    name = "02_SINTETIKE_zbatimi_buxhetit_jan-gus_2026.xlsx"
    monkeypatch.setattr(reader, "MAX_ROWS", 50)
    with pytest.raises(IngestError):
        reader.read_table((SAMPLES / name).read_bytes(), name)


# ---------------------------------------------------------------- export and indicator fields


@pytest.mark.parametrize("locale", ["sq", "en"])
def test_export_stock_note_comes_from_the_registry(client, locale):
    client.post("/api/v1/demo/reset")
    r = client.get(f"/api/v1/export/core_kpi.xlsx?locale={locale}")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    notes = " ".join(
        str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value
    )
    stock = [p.code for p in reg.load_pack().passports if p.period_kind == "point"]
    assert stock == ["REQ-03", "HR-01"]
    assert "(REQ-03, HR-01)" in notes
    assert "HR-02)" not in notes


def test_indicator_summary_says_how_to_read_series_and_smp(client):
    client.post("/api/v1/demo/reset")
    board = client.get("/api/v1/indicators").json()
    by_code = {i["code"]: i for i in board["indicators"]}
    assert by_code["FIN-01"]["smp_kind"] == "internal_view"
    assert by_code["FIN-01"]["smp_note"]["sq"].startswith("Pamje e brendshme")
    assert by_code["REQ-02"]["smp_kind"] == "direct"
    for code in ("FIN-01", "FIN-02", "REV-01", "REV-02"):
        assert by_code[code]["series_kind"] == "ytd_running"
    assert by_code["REQ-02"]["series_kind"] == "monthly"
    assert by_code["REQ-03"]["period_kind"] == "point"
