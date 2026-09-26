"""Ingest pipeline on the committed synthetic samples: preview, commit, lineage, recipes, reset."""

import csv
import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import catalog as c
from app.indicators import registry as reg
from app.ingest import pipeline
from app.ingest.errors import IngestError

SAMPLES = Path(__file__).resolve().parents[1] / "samples"
REQUESTS = "01_SINTETIKE_kerkesat_qytetare_jan-gus_2026.xlsx"
BUDGET = "02_SINTETIKE_zbatimi_buxhetit_jan-gus_2026.xlsx"
POPULATION = "ref_SINTETIKE_popullsia_njesite.csv"
WASTE = "zarfi-1_SINTETIKE_pastrimi_mbetjet_2026.csv"
REVENUE = "zarfi-2_SINTETIKE_taksat_tarifat_arketimi_2026.xlsx"
STAFF = "zarfi-3_SINTETIKE_burimet_njerezore_2026.xlsx"
DRIFT = "drift_SINTETIKE_pastrimi_mbetjet_2026_v2.csv"

# name -> (dataset, header_row, data_rows, excluded {row_no: reason}, pii columns, multiplier)
EXPECTED = {
    REQUESTS: (
        "requests",
        3,
        2400,
        {1: "title", 2: "title", 2404: "blank", 2405: "total_row"},
        {"Emri i kërkuesit": "name", "Nr. telefoni": "phone"},
        1,
    ),
    BUDGET: (
        "budget",
        3,
        144,
        {1: "title", 2: "title", 148: "blank", 149: "total_row"},
        {},
        1000,
    ),
    POPULATION: ("population", 2, 26, {1: "title"}, {}, 1),
    WASTE: ("waste", 2, 104, {1: "title", 107: "blank", 108: "total_row"}, {}, 1),
    REVENUE: (
        "revenue",
        3,
        96,
        {1: "title", 2: "title", 100: "blank", 101: "total_row"},
        {},
        1000,
    ),
    STAFF: (
        "staff",
        3,
        13,
        {1: "title", 2: "title", 17: "blank", 18: "total_row"},
        {"Përgjegjësi i drejtorisë": "name"},
        1,
    ),
    DRIFT: ("waste", 2, 104, {1: "title", 107: "blank", 108: "total_row"}, {}, 1),
}


def data(name: str) -> bytes:
    return (SAMPLES / name).read_bytes()


def mapping_of(preview: dict) -> list[dict]:
    return [{"column": m["column"], "field": m["field"]} for m in preview["mapping"]]


@pytest.mark.parametrize("name", list(EXPECTED))
def test_every_sample_previews_correctly(db, name):
    dataset, header_row, data_rows, excluded, pii, mult = EXPECTED[name]
    p = pipeline.preview(data(name), name, con=db)

    assert p["dataset"]["key"] == dataset
    assert p["dataset"]["confidence"] >= 0.6
    assert p["dataset_candidates"][0]["key"] == dataset
    assert p["header_row"] == header_row
    assert p["data_rows"] == data_rows
    assert {e["row_no"]: e["reason"] for e in p["excluded_rows"]} == excluded
    assert {col["name"]: col["pii"] for col in p["columns"] if col["dropped"]} == pii
    assert p["unit_multiplier"] == mult
    assert (p["unit_note"] is not None) == (mult > 1)
    assert p["synthetic"] is True
    assert len(p["file_hash"]) == 64 and p["size_bytes"] == len(data(name))
    assert p["sheet"] == (None if name.endswith(".csv") else p["sheet"])

    # every step reports a real result with both languages and a duration
    codes = [s["code"] for s in p["steps"]]
    assert codes == ["read", "header", "exclude", "pii", "profile", "dataset", "recipe", "mapping"]
    for s in p["steps"]:
        assert s["status"] in ("ok", "warn", "info")
        assert s["message"]["sq"] and s["message"]["en"]
        assert isinstance(s["ms"], int)
    pii_step = next(s for s in p["steps"] if s["code"] == "pii")["message"]
    if len(pii) == 1:  # singular agreement in both languages
        assert pii_step["sq"].startswith("1 kolonë personale u hoq para")
        assert pii_step["en"].startswith("1 personal column removed")
    elif pii:
        assert pii_step["sq"].startswith(f"{len(pii)} kolona personale u hoqën")

    # rules mode (no key on this machine): every required field mapped, all amber
    fields = {m["field"] for m in p["mapping"] if m["field"]}
    assert set(c.get_dataset(dataset).required_fields) <= fields
    assert all(m["confidence"] <= 0.6 and m["source"] == "rules" for m in p["mapping"])
    assert p["llm"]["used"] is False and p["llm"]["error"] == "llm_unavailable"
    # dropped columns never appear in the mapping and have no samples
    dropped = {col["name"] for col in p["columns"] if col["dropped"]}
    assert not dropped & {m["column"] for m in p["mapping"]}
    assert all(col["samples"] == [] for col in p["columns"] if col["dropped"])
    assert all(len(col["samples"]) <= 5 for col in p["columns"])


def test_personal_values_never_leave_the_gate(db):
    p = pipeline.preview(data(REQUESTS), REQUESTS, con=db)
    state = pipeline.get_preview(p["preview_id"])
    idx = [col["index"] for col in p["columns"] if col["dropped"]]
    assert idx == [2, 3]
    assert all(row[j] is None for _, row in state.layout.data for j in idx)
    text = json.dumps(p, ensure_ascii=False)
    assert "+355" not in text
    assert "Lindita Ç." not in text  # first requester in the file


def test_samples_are_masked(db):
    p = pipeline.preview(data(BUDGET), BUDGET, con=db)
    samples = {col["name"]: col["samples"] for col in p["columns"]}
    assert samples["Kodi i programit"][0] == "01110"  # codes keep their leading zero
    rev = pipeline.preview(data(WASTE), WASTE, con=db)
    tonnes = {col["name"]: col for col in rev["columns"]}["Sasia (ton)"]
    assert tonnes["inferred_type"] == "float"
    assert "2.562,4" in tonnes["samples"]


def test_commit_reconciles_and_applies_units(db):
    for name in (BUDGET, REVENUE, WASTE, STAFF, REQUESTS):
        p = pipeline.preview(data(name), name, con=db)
        r = pipeline.commit(p["preview_id"], p["dataset"]["key"], mapping_of(p), con=db)
        assert r["reconciliation"], name
        assert all(x["ok"] for x in r["reconciliation"]), (name, r["reconciliation"])
        assert r["rows_read"] == r["rows_loaded"] + sum(x["count"] for x in r["rows_excluded"])
        assert r["synthetic"] is True
        assert r["dataset_name"] == c.get_dataset(p["dataset"]["key"]).name
    plan, actual = db.execute(
        "SELECT sum(planned_lek), sum(actual_lek) FROM budget_line"
    ).fetchone()
    assert plan == pytest.approx(2_526_229.1 * 1000)
    assert actual == pytest.approx(1_996_945.5 * 1000)
    collected = db.execute("SELECT sum(collected_lek) FROM revenue").fetchone()[0]
    assert collected == pytest.approx(466_980.6 * 1000)
    assert db.execute("SELECT sum(tonnes) FROM waste_collection").fetchone()[0] == pytest.approx(
        26_828.9
    )
    assert db.execute("SELECT sum(headcount) FROM staff").fetchone()[0] == 1501
    # staff as_of comes from the header "Numri i punonjësve (31.08.2026)"
    assert {
        r[0].isoformat() for r in db.execute("SELECT DISTINCT as_of FROM staff").fetchall()
    } == {"2026-08-31"}
    # the programme code keeps its leading zero, line types are canonical
    assert (
        db.execute("SELECT count(*) FROM budget_line WHERE programme_code = '05100'").fetchone()[0]
        == 16
    )
    assert {r[0] for r in db.execute("SELECT DISTINCT line_type FROM budget_line").fetchall()} == {
        "current",
        "capital",
    }


def test_requests_receipt_reconciles_row_count(db):
    r = pipeline.ingest_path(SAMPLES / REQUESTS, con=db)
    assert r["rows_read"] == 2404 and r["rows_loaded"] == 2400
    assert r["pii_dropped"] == ["Emri i kërkuesit", "Nr. telefoni"]
    (rc,) = r["reconciliation"]
    assert rc["field"] == "row_count" and rc["file_total"] == 2400 and rc["ok"]
    assert [u["code"] for u in r["indicators_unlocked"]] == ["REQ-01", "REQ-02", "REQ-03", "REQ-04"]
    assert r["coverage"] == {"computable": 4, "total": 13}
    assert r["llm"] == {"used": False, "model": None, "latency_ms": None, "cost_usd": None}
    # personal values are not stored anywhere in the warehouse
    cols = [x[0] for x in db.execute("DESCRIBE request").fetchall()]
    assert not {"name", "phone"} & set(cols)
    stored = db.execute("SELECT receipt_json, mapping_json FROM source").fetchone()
    assert "+355" not in stored[0] and "Emri i kërkuesit" not in stored[1]


def test_row_no_is_the_original_file_row(db):
    pipeline.ingest_path(SAMPLES / REQUESTS, con=db)
    pipeline.ingest_path(SAMPLES / WASTE, con=db)

    ws = load_workbook(SAMPLES / REQUESTS, data_only=True).active
    rows = db.execute("SELECT row_no, request_id FROM request ORDER BY row_no").fetchall()
    assert rows[0] == (4, "KQ-2026-00001") and rows[-1] == (2403, "KQ-2026-02400")
    for row_no, request_id in rows[::97]:
        assert ws.cell(row=row_no, column=1).value == request_id

    lines = list(csv.reader(io.StringIO(data(WASTE).decode("utf-8-sig")), delimiter=";"))
    got = db.execute(
        "SELECT row_no, strftime(month, '%Y-%m'), admin_unit, tonnes FROM waste_collection"
    ).fetchall()
    for row_no, month, unit, tonnes in got:
        raw = lines[row_no - 1]
        assert raw[0] == month
        assert c.normalize_admin_unit(raw[1]) == unit
        assert c.parse_number(raw[2]) == pytest.approx(tonnes)


def test_admin_unit_variants_are_normalised(db):
    r = pipeline.ingest_path(SAMPLES / REQUESTS, con=db)
    step = next(s for s in r["steps"] if s["code"] == "admin_units")
    assert "Shirgjani" in step["message"]["sq"] or "Tregani" in step["message"]["sq"]
    pipeline.ingest_path(SAMPLES / WASTE, con=db)
    for table in ("request", "waste_collection"):
        units = {u for (u,) in db.execute(f"SELECT DISTINCT admin_unit FROM {table}").fetchall()}
        assert units <= set(c.ADMIN_UNITS), table
    assert (
        len({u for (u,) in db.execute("SELECT admin_unit FROM waste_collection").fetchall()}) == 13
    )


def test_recipe_is_reused_on_the_second_load(db):
    first = pipeline.ingest_path(SAMPLES / WASTE, save_recipe=True, con=db)
    assert first["recipe"]["saved"] is True and first["recipe"]["reused"] is False
    rid = first["recipe"]["recipe_id"]

    p = pipeline.preview(data(WASTE), WASTE, con=db)
    assert p["recipe"]["hit"] is True and p["recipe"]["recipe_id"] == rid
    assert p["recipe"]["drift"] is None
    assert all(m["source"] == "recipe" and m["confidence"] == 1.0 for m in p["mapping"])
    assert p["llm"]["used"] is False and p["llm"]["error"] is None
    second = pipeline.commit(p["preview_id"], "waste", mapping_of(p), save_recipe=True, con=db)
    assert second["recipe"] == {"saved": False, "reused": True, "recipe_id": rid}
    # the same export loaded twice replaces the first load: no double counting
    assert db.execute(
        "SELECT count(*), count(DISTINCT source_id) FROM waste_collection"
    ).fetchone() == (
        104,
        1,
    )
    assert second["superseded"][0]["rows"] == 104
    assert [s["source_id"] for s in pipeline.list_sources(db)] == [second["source_id"]]


def test_format_drift_is_detected_and_asks(db):
    pipeline.ingest_path(SAMPLES / WASTE, save_recipe=True, con=db)
    p = pipeline.preview(data(DRIFT), DRIFT, con=db)
    assert p["recipe"]["hit"] is False
    assert p["recipe"]["drift"] == {
        "renamed": ["Sasia (ton) → Tonazhi"],
        "missing": [],
        "added": ["Operatori"],
    }
    recipe_step = next(s for s in p["steps"] if s["code"] == "recipe")
    assert recipe_step["status"] == "warn"
    q = p["question"]
    assert q["column"] == "Tonazhi"
    assert [o["field"] for o in q["options"]] == ["tonnes", None]
    by_col = {m["column"]: m for m in p["mapping"]}
    assert by_col["Muaji"]["source"] == "recipe" and by_col["Muaji"]["confidence"] == 1.0
    assert by_col["Tonazhi"]["source"] != "recipe" and by_col["Tonazhi"]["confidence"] < 0.8
    assert by_col["Operatori"]["field"] is None
    r = pipeline.commit(p["preview_id"], "waste", mapping_of(p), save_recipe=True, con=db)
    assert r["recipe"]["saved"] is True
    assert all(x["ok"] for x in r["reconciliation"])


def test_reset_gives_6_of_13_and_the_envelopes_13(db):
    out = pipeline.reset_demo(db)
    assert out["ok"] is True
    assert out["coverage"] == {"computable": 6, "total": 13}
    assert {x["dataset"] for x in out["loaded"]} == {"requests", "budget", "population"}
    assert all(x["reconciled"] for x in out["loaded"])
    assert db.execute("SELECT count(*) FROM mapping_recipe").fetchone()[0] == 0

    unlocked = {}
    for name in (WASTE, REVENUE, STAFF):
        r = pipeline.ingest_path(SAMPLES / name, con=db)
        unlocked[name] = [u["code"] for u in r["indicators_unlocked"]]
    assert unlocked == {
        WASTE: ["WST-01", "WST-02", "WST-03"],
        REVENUE: ["REV-01", "REV-02"],
        STAFF: ["HR-01", "HR-02"],
    }
    assert reg.coverage(db)["computable"] == 13
    for p in reg.load_pack().passports:
        assert reg.compute(p, db).value is not None, p.code

    # a reset keeps the LLM call log but drops data and recipes again
    pipeline.ingest_path(SAMPLES / WASTE, save_recipe=True, con=db)
    assert pipeline.reset_demo(db)["coverage"]["computable"] == 6
    assert db.execute("SELECT count(*) FROM mapping_recipe").fetchone()[0] == 0


def test_commit_validates_the_mapping(db):
    p = pipeline.preview(data(REQUESTS), REQUESTS, con=db)
    pid = p["preview_id"]
    base = mapping_of(p)

    def fails(mapping, code, dataset="requests"):
        with pytest.raises(IngestError) as exc:
            pipeline.commit(pid, dataset, mapping, con=db)
        assert exc.value.code == code

    fails([m for m in base if m["field"] != "status"], "missing_required_fields")
    fails([*base, {"column": "Emri i kërkuesit", "field": "department"}], "personal_column")
    fails([*base, {"column": "Nope", "field": None}], "unknown_column")
    fails(
        [{**m, "field": "tonnes"} if m["field"] == "category" else m for m in base], "unknown_field"
    )
    dup = [{**m, "field": "category"} if m["field"] == "channel" else m for m in base]
    fails(dup, "duplicate_field")
    fails(base, "unknown_dataset", dataset="nope")
    # nothing was written by the failed attempts; the preview is still usable
    assert db.execute("SELECT count(*) FROM request").fetchone()[0] == 0
    r = pipeline.commit(pid, "requests", base, con=db)
    assert r["rows_loaded"] == 2400
    with pytest.raises(IngestError) as exc:
        pipeline.commit(pid, "requests", base, con=db)
    assert exc.value.code == "preview_not_found" and exc.value.status_code == 404


def test_remap_to_another_dataset(db):
    p = pipeline.preview(data(REVENUE), REVENUE, con=db)
    assert p["dataset"]["key"] == "revenue"
    q = pipeline.remap(p["preview_id"], "budget", con=db)
    assert q["preview_id"] == p["preview_id"]
    assert q["dataset"]["key"] == "budget"
    assert {m["field"] for m in q["mapping"]} >= {"month", "planned_lek"}
    assert any("Fushat" in w["sq"] or "Fusha" in w["sq"] for w in q["warnings"])  # missing fields


def test_ingest_path_requires_a_recognisable_dataset(db, tmp_path):
    f = tmp_path / "shenime.csv"
    f.write_text("Titulli\nabc;def\nx;y\n", encoding="utf-8")
    with pytest.raises(IngestError) as exc:
        pipeline.ingest_path(f, con=db)
    assert exc.value.code == "dataset_not_recognised"


def test_sources_list_has_receipts_and_mapping(db):
    r = pipeline.ingest_path(SAMPLES / REVENUE, con=db)
    (src,) = pipeline.list_sources(db)
    assert src["source_id"] == r["source_id"]
    assert src["mapping"] == r["mapping"]
    for key in (
        "filename",
        "file_hash",
        "dataset",
        "dataset_name",
        "synthetic",
        "rows_read",
        "rows_loaded",
        "rows_excluded",
        "reconciliation",
        "pii_dropped",
        "unit_multiplier",
        "llm",
        "recipe",
        "indicators_unlocked",
        "coverage",
        "loaded_at",
        "duration_ms",
    ):
        assert key in src, key
    assert src["unit_multiplier"] == 1000
    assert [u["code"] for u in src["indicators_unlocked"]] == ["REV-01"]  # budget not loaded
