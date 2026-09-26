"""AI column mapping through a fake SDK: what is sent, how answers are checked, fallback."""

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.ingest import pipeline
from app.llm import client as llm_client
from app.llm.client import MODEL_FAST, LLMClient

SAMPLES = Path(__file__).resolve().parents[1] / "samples"
REQUESTS = "01_SINTETIKE_kerkesat_qytetare_jan-gus_2026.xlsx"
WASTE = "zarfi-1_SINTETIKE_pastrimi_mbetjet_2026.csv"


class FakeMessages:
    def __init__(self, answer: dict | None = None, fail: bool = False):
        self.answer = answer
        self.fail = fail
        self.calls: list[dict] = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        if self.fail:
            raise RuntimeError("network down")
        return SimpleNamespace(
            content=[SimpleNamespace(type="text", text=json.dumps(self.answer))],
            stop_reason="end_turn",
            usage=SimpleNamespace(
                input_tokens=1800,
                output_tokens=400,
                cache_creation_input_tokens=0,
                cache_read_input_tokens=0,
            ),
        )


def install(db, **kw) -> FakeMessages:
    messages = FakeMessages(**kw)
    sdk = SimpleNamespace(messages=messages)
    llm_client._client = LLMClient(api_key="test", budget_usd=5.0, sdk=sdk, db=lambda: db)
    return messages


def item(column, field, confidence=0.95):
    return {
        "column": column,
        "field": field,
        "confidence": confidence,
        "reason_sq": f"Koka '{column}' përputhet.",
        "reason_en": f"Header '{column}' fits.",
    }


REQUESTS_ANSWER = {
    "columns": [
        item("Nr.", "request_id"),
        item("Data e regjistrimit", "created_at", 0.97),
        item("Kategoria", "category"),
        item("Drejtoria përgjegjëse", "department", 0.9),
        item("Njësia administrative", "admin_unit"),
        item("Kanali", "sla_days", 0.9),  # wrong type: text column → int field
        item("Statusi", "status"),
        item("Data e mbylljes", "created_at", 0.7),  # duplicate field, less confident
        item("Afati (ditë)", "no_such_field", 0.8),  # unknown field
    ],
    "question": {
        "column": "Data e mbylljes",
        "text_sq": "A është kjo data e mbylljes?",
        "text_en": "Is this the closing date?",
        "options": ["closed_at", None, "bogus"],
    },
}


def test_ai_mapping_is_used_and_checked_by_code(db):
    messages = install(db, answer=REQUESTS_ANSWER)
    p = pipeline.preview((SAMPLES / REQUESTS).read_bytes(), REQUESTS, con=db)

    assert len(messages.calls) == 1
    call = messages.calls[0]
    assert call["model"] == MODEL_FAST
    assert call["output_config"]["format"]["type"] == "json_schema"

    llm = p["llm"]
    assert llm["used"] is True and llm["model"] == MODEL_FAST and llm["error"] is None
    assert llm["cost_usd"] == pytest.approx((1800 * 1.0 + 400 * 5.0) / 1e6)
    assert llm["sent"] == {"headers": 9, "samples_per_column": 5, "rows_sent": 0}

    by_col = {m["column"]: m for m in p["mapping"]}
    assert by_col["Nr."] == {**by_col["Nr."], "field": "request_id", "source": "ai"}
    assert by_col["Nr."]["confidence"] == 0.95
    assert by_col["Data e regjistrimit"]["field"] == "created_at"
    # code checks the AI: type mismatch capped at amber, unknown field dropped, duplicate removed
    assert by_col["Kanali"]["field"] == "sla_days" and by_col["Kanali"]["confidence"] <= 0.5
    assert "Code check" in by_col["Kanali"]["reason"]["en"]
    assert by_col["Afati (ditë)"]["field"] is None
    assert by_col["Data e mbylljes"]["field"] is None
    assert p["question"]["column"] == "Data e mbylljes"
    assert [o["field"] for o in p["question"]["options"]] == ["closed_at", None]
    step = next(s for s in p["steps"] if s["code"] == "mapping")
    assert step["status"] == "ok" and "Claude Haiku" in step["message"]["en"]


def test_personal_data_is_never_sent_to_the_model(db):
    messages = install(db, answer=REQUESTS_ANSWER)
    pipeline.preview((SAMPLES / REQUESTS).read_bytes(), REQUESTS, con=db)
    call = messages.calls[0]
    sent_text = call["system"] + json.dumps(call["messages"], ensure_ascii=False)

    ws = load_workbook(SAMPLES / REQUESTS, data_only=True).active
    names = {r[2] for r in ws.iter_rows(min_row=4, max_row=2403, values_only=True)}
    phones = {r[3] for r in ws.iter_rows(min_row=4, max_row=2403, values_only=True)}
    assert not any(n in sent_text for n in names)
    assert not any(ph in sent_text for ph in phones)
    assert "+355" not in sent_text
    assert "Emri i kërkuesit" not in sent_text and "Nr. telefoni" not in sent_text
    payload = json.loads(call["messages"][0]["content"].split("\n\n", 1)[1])
    assert len(payload["columns"]) == 9
    assert all(len(col["samples"]) <= 5 for col in payload["columns"])

    # the call log records what left the building: headers and samples, never rows
    row = db.execute("SELECT purpose, ok, sent_json FROM llm_call").fetchone()
    assert row[0] == "ingest_mapping" and row[1] is True
    sent = json.loads(row[2])
    assert sent["rows_sent"] == 0 and sent["headers"] == 9 and sent["pii_columns_withheld"] == 2


def test_failed_call_falls_back_to_rules(db):
    messages = install(db, fail=True)
    p = pipeline.preview((SAMPLES / WASTE).read_bytes(), WASTE, con=db)
    assert len(messages.calls) == 1
    assert p["llm"]["used"] is True
    assert p["llm"]["error"] == "llm_error:unexpected"
    assert all(m["source"] == "rules" and m["confidence"] <= 0.6 for m in p["mapping"])
    assert {m["field"] for m in p["mapping"]} >= {"month", "admin_unit", "tonnes"}
    step = next(s for s in p["steps"] if s["code"] == "mapping")
    assert step["status"] == "warn"
    # the receipt says the AI was attempted
    r = pipeline.commit(
        p["preview_id"],
        "waste",
        [{"column": m["column"], "field": m["field"]} for m in p["mapping"]],
        con=db,
    )
    assert r["llm"]["used"] is True and r["llm"]["model"] == MODEL_FAST
    assert all(x["ok"] for x in r["reconciliation"])


def test_columns_the_model_skips_fall_back_to_rules(db):
    answer = {
        "columns": [item("Muaji", "month"), item("Sasia (ton)", "tonnes", 0.92)],
        "question": {"column": None, "text_sq": "", "text_en": "", "options": []},
    }
    install(db, answer=answer)
    p = pipeline.preview((SAMPLES / WASTE).read_bytes(), WASTE, con=db)
    by_col = {m["column"]: m for m in p["mapping"]}
    assert by_col["Muaji"]["source"] == "ai" and by_col["Sasia (ton)"]["confidence"] == 0.92
    assert by_col["Njësia administrative"]["source"] == "rules"
    assert by_col["Njësia administrative"]["field"] == "admin_unit"
    assert p["question"] is None


def test_recipe_hit_skips_the_model(db):
    pipeline.ingest_path(SAMPLES / WASTE, save_recipe=True, con=db)
    messages = install(db, answer={"columns": [], "question": None})
    p = pipeline.preview((SAMPLES / WASTE).read_bytes(), WASTE, con=db)
    assert messages.calls == []
    assert p["recipe"]["hit"] is True and p["llm"]["used"] is False


def test_ingest_path_uses_rules_unless_asked(db):
    messages = install(db, answer={"columns": [], "question": None})
    r = pipeline.ingest_path(SAMPLES / WASTE, con=db)
    assert messages.calls == [] and r["llm"]["used"] is False
