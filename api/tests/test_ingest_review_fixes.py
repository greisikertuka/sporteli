"""Regression tests for the ingest review findings: composite supersede keys, stale
reconciliation, rounding-level reconciliation tolerance and the atomic demo reset."""

import shutil
import threading
from pathlib import Path

from openpyxl import load_workbook

from app.indicators import signals as sig
from app.ingest import pipeline

SAMPLES = Path(__file__).resolve().parents[1] / "samples"
BUDGET = "02_SINTETIKE_zbatimi_buxhetit_jan-gus_2026.xlsx"
POPULATION = "ref_SINTETIKE_popullsia_njesite.csv"
WASTE = "zarfi-1_SINTETIKE_pastrimi_mbetjet_2026.csv"


def _write(tmp_path: Path, name: str, text: str) -> Path:
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


def test_partial_waste_export_replaces_only_its_own_unit_and_month(db, tmp_path):
    first = pipeline.ingest_path(SAMPLES / WASTE, "waste", con=db)
    july = "SELECT count(*), round(sum(tonnes), 1) FROM waste_collection WHERE month = '2026-07-01'"
    n_july, t_july = db.execute(july).fetchone()
    total = db.execute("SELECT round(sum(tonnes), 1) FROM waste_collection").fetchone()[0]
    gjinar_july = db.execute(
        "SELECT tonnes FROM waste_collection WHERE month = '2026-07-01' AND admin_unit = 'Gjinar'"
    ).fetchone()[0]

    fix = _write(
        tmp_path,
        "korrigjim_SINTETIKE_mbetjet.csv",
        "Muaji;Njësia administrative;Sasia (ton)\n2026-07;Gjinar;21,4\n",
    )
    second = pipeline.ingest_path(fix, "waste", con=db)

    # only the (2026-07, Gjinar) row was replaced; the other 12 units of July stay
    assert second["superseded"] == [
        {
            "source_id": first["source_id"],
            "filename": WASTE,
            "rows": 1,
            "source_removed": False,
            "rows_remaining": 103,
            "key": ["month", "admin_unit"],
        }
    ]
    assert db.execute(july).fetchone() == (n_july, round(t_july - gjinar_july + 21.4, 1))
    assert db.execute("SELECT count(*) FROM waste_collection").fetchone()[0] == 104
    assert db.execute("SELECT round(sum(tonnes), 1) FROM waste_collection").fetchone()[0] == round(
        total - gjinar_july + 21.4, 1
    )

    # the earlier file's reconciliation no longer describes the stored rows: marked stale,
    # and parts_vs_total neither passes nor warns on it
    raw = db.execute(
        "SELECT reconciliation_json FROM source WHERE id = ?", [first["source_id"]]
    ).fetchone()[0]
    recons = sig.parse_reconciliation(WASTE, raw)
    assert recons and all(r.stale for r in recons)
    assert sig.parts_vs_total_signals(recons) == []
    listed = {s["source_id"]: s for s in pipeline.list_sources(db)}
    assert listed[first["source_id"]]["rows_current"] == 103
    assert listed[first["source_id"]]["superseded_by"] == [
        {"source_id": second["source_id"], "rows": 1}
    ]
    assert all(r["stale"] for r in listed[first["source_id"]]["reconciliation"])


def test_partial_population_export_keeps_other_units(db, tmp_path):
    pipeline.ingest_path(SAMPLES / POPULATION, "population", con=db)
    before = db.execute(
        "SELECT count(*), sum(residents) FROM population WHERE basis = 'census_2023'"
    ).fetchone()
    fix = _write(
        tmp_path,
        "pop_fix_SINTETIKE.csv",
        "Njësia administrative,Baza,Banorë\nElbasan,census_2023,83700\nGjinar,census_2023,1300\n",
    )
    out = pipeline.ingest_path(fix, "population", con=db)
    assert sum(s["rows"] for s in out["superseded"]) == 2
    after = db.execute(
        "SELECT count(*), sum(residents) FROM population WHERE basis = 'census_2023'"
    ).fetchone()
    assert after[0] == before[0]
    assert db.execute("SELECT count(*) FROM population").fetchone()[0] == 26


def test_same_export_twice_still_replaces_everything(db):
    pipeline.ingest_path(SAMPLES / WASTE, "waste", con=db)
    again = pipeline.ingest_path(SAMPLES / WASTE, "waste", con=db)
    assert again["superseded"][0]["rows"] == 104
    assert again["superseded"][0]["source_removed"] is True
    assert db.execute(
        "SELECT count(*), count(DISTINCT source_id) FROM waste_collection"
    ).fetchone() == (104, 1)


def test_reconciliation_fails_when_invalid_rows_carry_money(db, tmp_path):
    path = tmp_path / "buxheti_SINTETIKE_me_gabime.xlsx"
    shutil.copy(SAMPLES / BUDGET, path)
    wb = load_workbook(path)
    ws = wb.active
    dropped_plan = float(ws.cell(5, 5).value) + float(ws.cell(30, 5).value)
    dropped_actual = float(ws.cell(5, 6).value)
    ws.cell(5, 1).value = "??"
    ws.cell(30, 6).value = "n/a"
    wb.save(path)

    receipt = pipeline.ingest_path(path, "budget", con=db)
    assert receipt["rows_loaded"] == 142
    rec = {r["field"]: r for r in receipt["reconciliation"]}
    assert rec["planned_lek"]["ok"] is False
    assert rec["actual_lek"]["ok"] is False
    assert abs(rec["planned_lek"]["excluded_sum"] - dropped_plan * 1000) < 1
    assert abs(rec["actual_lek"]["excluded_sum"] - dropped_actual * 1000) < 1
    assert "not loaded" in rec["planned_lek"]["note"]["en"]
    step = next(s for s in receipt["steps"] if s["code"] == "reconcile")
    assert step["status"] == "warn"
    assert "0 of 2" in step["message"]["en"]


def test_reconciliation_passes_to_rounding_only(db):
    receipt = pipeline.ingest_path(SAMPLES / BUDGET, "budget", con=db)
    for r in receipt["reconciliation"]:
        assert r["ok"] is True
        assert r["excluded_sum"] == 0
        # half a unit of the file's precision (values in thousands of lek), plus float epsilon
        assert 500 <= r["tolerance"] < 510


def test_demo_reset_is_atomic_for_concurrent_readers(db):
    from app.warehouse.db import new_cursor

    pipeline.reset_demo(db)
    errors: list[BaseException] = []
    counts: list[int] = []
    stop = threading.Event()

    def reader() -> None:
        cur = new_cursor()
        try:
            while not stop.is_set():
                try:
                    n = cur.execute(
                        "SELECT (SELECT count(*) FROM request) + (SELECT count(*) FROM budget_line)"
                    ).fetchone()[0]
                    counts.append(int(n))
                except Exception as exc:  # noqa: BLE001 - any error is a failure here
                    errors.append(exc)
        finally:
            cur.close()

    threads = [threading.Thread(target=reader) for _ in range(3)]
    for th in threads:
        th.start()
    try:
        for _ in range(2):
            writer = new_cursor()
            try:
                pipeline.reset_demo(writer)
            finally:
                writer.close()
    finally:
        stop.set()
        for th in threads:
            th.join()
    assert errors == []
    # readers only ever saw the complete preload (2400 requests + 144 budget lines)
    assert counts and set(counts) == {2544}
