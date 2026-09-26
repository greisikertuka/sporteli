"""End to end: the real synthetic sample files through the ingest pipeline into the board.

Skipped when ``app.ingest.pipeline.ingest_path`` is not available (the ingest module is
built by another agent). Starts from the demo's preloaded files (6/13), then loads the three
envelopes and expects 13/13, the planted signals, and the real filenames in the export.
"""

import csv
import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

try:  # pragma: no cover - depends on the ingest module being present
    from app.ingest.pipeline import ingest_path
except Exception:  # noqa: BLE001
    ingest_path = None

SAMPLES = Path(__file__).resolve().parents[1] / "samples"

pytestmark = pytest.mark.skipif(
    ingest_path is None, reason="app.ingest.pipeline.ingest_path is not available"
)


def manifest() -> list[dict]:
    return json.loads((SAMPLES / "manifest.json").read_text(encoding="utf-8"))["files"]


def board(client) -> dict:
    res = client.get("/api/v1/indicators")
    assert res.status_code == 200, res.text
    return res.json()


def test_real_samples_from_gap_to_proof(client):
    files = {f["name"]: f for f in manifest()}
    preload = [f for f in files.values() if f["preload"]]
    envelopes = sorted((f for f in files.values() if f["envelope"]), key=lambda f: f["envelope"])
    assert [f["dataset_hint"] for f in envelopes] == ["waste", "revenue", "staff"]

    for f in preload:
        ingest_path(SAMPLES / f["name"], f["dataset_hint"])
    b = board(client)
    assert b["coverage"]["computable"] == 6 and b["coverage"]["total"] == 13
    ind = {i["code"]: i for i in b["indicators"]}
    assert ind["REQ-01"]["value"] == 2400
    assert ind["REV-02"]["state"] == "missing"
    assert ind["REV-02"]["missing"][0]["dataset"] == "revenue"
    smp = client.get("/api/v1/coverage").json()
    assert smp["counts"]["computable"] == 0

    for f in envelopes:
        ingest_path(SAMPLES / f["name"], f["dataset_hint"])
    b = board(client)
    assert b["coverage"]["computable"] == 13 and b["as_of"] == "2026-08"
    ind = {i["code"]: i for i in b["indicators"]}
    assert all(i["value"] is not None for i in ind.values())
    assert all(s["synthetic"] for i in ind.values() for s in i["sources"])
    assert 35 <= ind["FIN-02"]["value"] <= 40 and ind["FIN-02"]["status"] == "off_track"
    assert 50 <= ind["REV-02"]["value"] <= 60

    fired = {(c, s["rule"], s["period"]) for c, i in ind.items() for s in i["signals"]}
    assert ("WST-01", "placeholder_value", "2026-07") in fired
    assert ("WST-03", "swing", "2026-05") in fired
    assert ("FIN-02", "off_target", "2026-08") in fired
    assert ("REQ-02", "swing", "2026-07") in fired
    assert not any(rule == "parts_vs_total" for _, rule, _ in fired)  # every file reconciles
    gjinar = next(s for s in ind["WST-01"]["signals"] if s["rule"] == "placeholder_value")
    assert gjinar["message"]["sq"].startswith("Gjinar: 0 ton në korrik 2026")

    passport = client.get("/api/v1/indicators/REV-02").json()
    assert {s["filename"] for s in passport["lineage"]} == {
        "zarfi-2_SINTETIKE_taksat_tarifat_arketimi_2026.xlsx",
        "02_SINTETIKE_zbatimi_buxhetit_jan-gus_2026.xlsx",
    }
    assert all(len(s["file_hash"]) == 64 for s in passport["lineage"])

    smp = client.get("/api/v1/coverage").json()
    assert smp["counts"]["computable"] == 5

    res = client.get("/api/v1/export/core_kpi.xlsx")
    assert "sportel-treguesit-2026-08.xlsx" in res.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(res.content))["Treguesit"]
    header = [c.value for c in ws[1]]
    burimi = {r[0].value: r[header.index("Burimi")].value for r in ws.iter_rows(min_row=2)}
    assert burimi["REQ-01"].startswith(
        "01_SINTETIKE_kerkesat_qytetare_jan-gus_2026.xlsx · rreshtat"
    )
    assert "zarfi-2_SINTETIKE_taksat_tarifat_arketimi_2026.xlsx" in burimi["REV-02"]

    text = client.get("/api/v1/open-data/indicators.csv").content.decode("utf-8-sig")
    assert len(list(csv.DictReader(io.StringIO(text)))) == 13
