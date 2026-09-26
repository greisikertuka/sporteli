"""GET /export/core_kpi.xlsx and GET /open-data/indicators.csv."""

import csv
import io
import json

from openpyxl import load_workbook

from app.warehouse.db import get_db
from tests.test_indicators import ALL, START, load_mini

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER_SQ = [
    "Kodi",
    "Treguesi",
    "Vlera",
    "Njësia",
    "Periudha",
    "Statusi",
    "Objektivi",
    "Burimi",
    "Baza",
    "Sinjalet",
]


def workbook(client, **params):
    res = client.get("/api/v1/export/core_kpi.xlsx", params=params)
    assert res.status_code == 200, res.text
    return res, load_workbook(io.BytesIO(res.content))


def rows_by_code(ws) -> dict[str, dict]:
    header = [c.value for c in ws[1]]
    out = {}
    for row in ws.iter_rows(min_row=2):
        values = {h: c for h, c in zip(header, row, strict=True)}
        out[values[header[0]].value] = values
    return out


def test_xlsx_download_headers(client):
    load_mini(get_db(), ALL)
    res, wb = workbook(client)
    assert res.headers["content-type"] == XLSX
    assert (
        res.headers["content-disposition"]
        == 'attachment; filename="sportel-treguesit-2026-06.xlsx"'
    )
    assert "Content-Disposition" in res.headers["access-control-expose-headers"]
    assert res.headers["x-synthetic-data"] == "true"
    assert wb.sheetnames == ["Treguesit", "Burimet", "Shënime"]
    assert wb.properties.creator == "Sportel"


def test_indicator_sheet_has_a_visible_source_column_and_comments(client):
    load_mini(get_db(), ALL)
    _, wb = workbook(client)
    ws = wb["Treguesit"]
    assert [c.value for c in ws[1]] == HEADER_SQ
    assert ws.freeze_panes == "C2"
    assert ws.auto_filter.ref == "A1:J14"
    assert ws.column_dimensions["H"].width >= 60  # Burimi is wide enough to read on a projector
    assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb.endswith("1F3A5F")
    rows = rows_by_code(ws)
    assert len(rows) == 13

    req = rows["REQ-01"]
    assert req["Vlera"].value == 24 and req["Vlera"].number_format == "#,##0"
    assert req["Njësia"].value == "kërkesa" and req["Periudha"].value == "janar – qershor 2026"
    assert req["Statusi"].value == "Pa objektiv" and req["Objektivi"].value == "—"
    burimi = req["Burimi"].value
    assert burimi.startswith("01_SINTETIKE_kerkesat.xlsx · rreshtat 4–27 · pasaporta v0.1.0 · ")
    assert "llogaritur më" in burimi and burimi.endswith("UTC")
    comment = req["Vlera"].comment
    assert comment is not None and comment.author == "Sportel"
    assert "01_SINTETIKE_kerkesat.xlsx" in comment.text
    assert "Rreshtat: 4–27 (24 rreshta)" in comment.text
    assert "SHA-256: sha-s-req" in comment.text
    assert "REQ-01 v0.1.0 · formula draft — në pritje të validimit" in comment.text
    assert "TË DHËNA SINTETIKE" in comment.text

    wst3 = rows["WST-03"]["Burimi"].value
    assert "zarfi-1_SINTETIKE_mbetjet.csv · rreshtat 3–14 + 02_SINTETIKE_buxheti.xlsx" in wst3
    assert "4–5, 8–9, 12–13 … (12 rreshta)" in wst3  # long range lists are shortened

    fin2 = rows["FIN-02"]
    assert fin2["Statusi"].value == "Jashtë objektivit" and fin2["Objektivi"].value == "≥ 70%"
    assert fin2["Statusi"].fill.fgColor.rgb.endswith("FDF0DB")
    assert "nën objektivin 70%" in fin2["Sinjalet"].value
    assert rows["REQ-04"]["Objektivi"].value == "≤ 10 ditë"
    assert rows["REQ-04"]["Statusi"].value == "Në objektiv"

    hr1 = rows["HR-01"]
    assert hr1["Baza"].value == "Censusi 2023"
    assert "ref_SINTETIKE_popullsia.csv" in hr1["Burimi"].value
    assert "Baza e popullsisë: Censusi 2023" in hr1["Vlera"].comment.text
    assert rows["WST-01"]["Sinjalet"].value.startswith("• Gjinar: 0 ton në mars 2026")
    assert rows["REQ-01"]["Sinjalet"].value == "—"


def test_missing_indicators_are_listed_with_their_owner(client):
    load_mini(get_db(), START)
    _, wb = workbook(client)
    rows = rows_by_code(wb["Treguesit"])
    rev1 = rows["REV-01"]
    assert rev1["Vlera"].value == "—" and rev1["Vlera"].comment is None
    assert rev1["Statusi"].value == (
        "Mungon: Taksat dhe tarifat — arkëtimi · Përgjegjës: Drejtoria e të Ardhurave Vendore"
    )
    assert rev1["Burimi"].value == "—" and rev1["Kodi"].font.italic
    assert rows["REQ-01"]["Burimi"].value.startswith("01_SINTETIKE_kerkesat.xlsx")


def test_sources_and_notes_sheets(client):
    load_mini(get_db(), ALL)
    recon = [
        {"field": "tonnes", "label": {"sq": "Sasia (ton)", "en": "Quantity (tonnes)"},
         "file_total": 700.0, "loaded_sum": 700.0, "ok": True, "note": None},
    ]  # fmt: skip
    get_db().execute(
        "UPDATE source SET reconciliation_json = ?, rows_excluded_json = ? WHERE id = 's-wst'",
        [json.dumps(recon), json.dumps([{"reason": "total_row", "count": 1, "label": {}}])],
    )
    _, wb = workbook(client)
    ws = wb["Burimet"]
    header = [c.value for c in ws[1]]
    assert header[:3] == ["Skedari", "Të dhënat", "SHA-256"]
    by_file = {
        r[0].value: {h: c.value for h, c in zip(header, r, strict=True)}
        for r in ws.iter_rows(min_row=2)
    }
    assert set(by_file) == {
        "01_SINTETIKE_kerkesat.xlsx",
        "02_SINTETIKE_buxheti.xlsx",
        "ref_SINTETIKE_popullsia.csv",
        "zarfi-1_SINTETIKE_mbetjet.csv",
        "zarfi-2_SINTETIKE_taksat.xlsx",
        "zarfi-3_SINTETIKE_bnj.xlsx",
    }
    waste = by_file["zarfi-1_SINTETIKE_mbetjet.csv"]
    assert waste["SHA-256"] == "sha-s-wst" and waste["Të dhënat"] == "Pastrimi dhe mbetjet"
    assert waste["Rreshta të ngarkuar"] == 12 and waste["Rreshta të përjashtuar"] == 1
    assert waste["Kontrolli i totalit"] == "✓ përputhet (Sasia (ton))"
    assert waste["Sintetike"] == "Po"
    assert waste["Treguesit"] == "WST-01, WST-02, WST-03"
    assert by_file["02_SINTETIKE_buxheti.xlsx"]["Treguesit"] == ("FIN-01, FIN-02, WST-03, REV-02")

    notes = "\n".join(
        str(c.value) for row in wb["Shënime"].iter_rows() for c in row if c.value is not None
    )
    assert "TË DHËNA SINTETIKE" in notes
    assert "Formulat janë draft — në pritje të validimit" in notes
    assert "AI nuk shkruan asnjë shifër" in notes
    assert "Censusi 2023" in notes
    assert "SMP-AL 2024 #47 · Pamje e brendshme" in notes  # FIN-01 reference with its note


def test_xlsx_in_english(client):
    load_mini(get_db(), ALL)
    res, wb = workbook(client, locale="en")
    assert wb.sheetnames == ["Indicators", "Sources", "Notes"]
    assert "sportel-indicators-2026-06.xlsx" in res.headers["content-disposition"]
    rows = rows_by_code(wb["Indicators"])
    assert rows["REQ-01"]["Source"].value.startswith("01_SINTETIKE_kerkesat.xlsx · rows 4–27 · ")
    assert rows["FIN-02"]["Status"].value == "Off track"
    assert client.get("/api/v1/export/core_kpi.xlsx", params={"locale": "de"}).status_code == 422


def test_xlsx_without_data(client):
    res, wb = workbook(client)
    assert "sportel-treguesit-pa-te-dhena.xlsx" in res.headers["content-disposition"]
    rows = rows_by_code(wb["Treguesit"])
    assert len(rows) == 13 and all(r["Vlera"].value == "—" for r in rows.values())
    assert wb["Burimet"].max_row == 1


# --------------------------------------------------------------------------------------------
# Open-data CSV
# --------------------------------------------------------------------------------------------


def test_open_data_csv(client):
    load_mini(get_db(), START)
    res = client.get("/api/v1/open-data/indicators.csv")
    assert res.status_code == 200
    assert res.headers["content-type"] == "text/csv; charset=utf-8"
    assert res.headers["x-license"] == "CC-BY-4.0"
    assert 'rel="license"' in res.headers["link"] and "by/4.0" in res.headers["link"]
    assert res.content.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM for Excel
    text = res.content.decode("utf-8-sig")
    assert "\r\n" in text
    rows = list(csv.DictReader(io.StringIO(text)))
    assert list(rows[0])[:10] == [
        "code",
        "name_sq",
        "name_en",
        "value",
        "unit",
        "period",
        "basis",
        "sources",
        "version",
        "formula_status",
    ]
    assert [r["code"] for r in rows] == ["REQ-01", "REQ-02", "REQ-03", "REQ-04", "FIN-01", "FIN-02"]
    req = rows[0]
    assert req["value"] == "24" and req["unit"] == "count" and req["period"] == "2026-01/2026-06"
    assert req["name_sq"] == "Kërkesa të pranuara" and req["version"] == "0.1.0"
    assert req["formula_status"] == "draft" and req["synthetic"] == "true"
    assert req["sources"] == "01_SINTETIKE_kerkesat.xlsx (sha256:sha-s-req)"
    assert rows[1]["value"] == "66.6667" and rows[1]["target"] == "90"
    assert rows[1]["status"] == "off_track"


def test_open_data_csv_states_the_basis(client):
    load_mini(get_db(), ALL)
    client.post("/api/v1/definitions/population_basis", json={"value": "civil_registry"})
    text = client.get("/api/v1/open-data/indicators.csv").content.decode("utf-8-sig")
    rows = {r["code"]: r for r in csv.DictReader(io.StringIO(text))}
    assert len(rows) == 13
    assert rows["WST-02"]["basis"] == "civil_registry" and rows["WST-02"]["value"] == "700"
    assert rows["HR-01"]["basis"] == "civil_registry" and rows["REQ-01"]["basis"] == ""
    assert " | " in rows["WST-02"]["sources"]
