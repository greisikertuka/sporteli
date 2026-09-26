"""The committed synthetic samples: manifest, messiness, reconciliation and planted patterns.

The files are read here with deliberately simple, file-specific code (independent of the
ingest pipeline), loaded into the canonical tables and run through every passport.
"""

import csv
import datetime as dt
import hashlib
import importlib.util
import io
import json
from collections import defaultdict
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import catalog as c
from app.indicators import registry as reg
from app.warehouse.db import insert_rows

SAMPLES = Path(__file__).resolve().parents[1] / "samples"
SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "generate_samples.py"

EXPECTED_FILES = {
    "01_SINTETIKE_kerkesat_qytetare_jan-gus_2026.xlsx": ("requests", None, True),
    "02_SINTETIKE_zbatimi_buxhetit_jan-gus_2026.xlsx": ("budget", None, True),
    "ref_SINTETIKE_popullsia_njesite.csv": ("population", None, True),
    "zarfi-1_SINTETIKE_pastrimi_mbetjet_2026.csv": ("waste", 1, False),
    "zarfi-2_SINTETIKE_taksat_tarifat_arketimi_2026.xlsx": ("revenue", 2, False),
    "zarfi-3_SINTETIKE_burimet_njerezore_2026.xlsx": ("staff", 3, False),
    "drift_SINTETIKE_pastrimi_mbetjet_2026_v2.csv": ("waste", None, False),
}


def manifest() -> dict:
    return json.loads((SAMPLES / "manifest.json").read_text(encoding="utf-8"))


def xlsx_rows(name: str) -> list[tuple]:
    ws = load_workbook(SAMPLES / name, data_only=True).active
    return list(ws.iter_rows(values_only=True))


def csv_rows(name: str, delimiter: str) -> list[list[str]]:
    text = (SAMPLES / name).read_bytes().decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def test_manifest_lists_exactly_the_contract_files():
    m = manifest()
    assert m["synthetic"] is True
    files = {f["name"]: f for f in m["files"]}
    assert set(files) == set(EXPECTED_FILES)
    for name, (hint, envelope, preload) in EXPECTED_FILES.items():
        f = files[name]
        assert (f["dataset_hint"], f["envelope"], f["preload"]) == (hint, envelope, preload)
        assert f["synthetic"] is True and f["label"]["sq"] and f["label"]["en"]
        assert "SINTETIKE" in name
        assert (SAMPLES / name).is_file()


def test_every_file_has_a_synthetic_title_row():
    for name in EXPECTED_FILES:
        if name.endswith(".xlsx"):
            first = xlsx_rows(name)[0][0]
        else:
            first = (SAMPLES / name).read_bytes().decode("utf-8-sig").splitlines()[0]
        assert "SINTETIKE" in first, name


def test_generator_is_deterministic_and_committed_files_are_current(tmp_path, capsys):
    spec = importlib.util.spec_from_file_location("generate_samples", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    mod.OUT = tmp_path
    mod.main()
    capsys.readouterr()
    for name in [*EXPECTED_FILES, "manifest.json"]:
        fresh = hashlib.sha256((tmp_path / name).read_bytes()).hexdigest()
        committed = hashlib.sha256((SAMPLES / name).read_bytes()).hexdigest()
        assert fresh == committed, f"{name} is stale: run scripts/generate_samples.py"


def test_requests_file_shape_and_personal_columns():
    rows = xlsx_rows("01_SINTETIKE_kerkesat_qytetare_jan-gus_2026.xlsx")
    header = rows[2]
    assert header[0] == "Nr." and "Emri i kërkuesit" in header and "Nr. telefoni" in header
    data = [r for r in rows[3:] if r[0] and str(r[0]).startswith("KQ-")]
    assert 2300 <= len(data) <= 2500
    assert rows[-1][0] == "TOTALI" and str(len(data)) in rows[-1][1]
    phones = {r[3] for r in data}
    assert all(p.startswith("+355 69 000 ") for p in phones)
    units = {r[6] for r in data}
    assert {"Shirgjani", "Tregani"} <= units
    assert all(c.match_admin_unit(u) for u in units)


# --------------------------------------------------------------------------------------------
# Independent loader → canonical tables → passports
# --------------------------------------------------------------------------------------------


def load_all(con) -> dict:
    totals: dict[str, dict] = {}

    # requests: header row 3, data until the blank row
    rows = xlsx_rows("01_SINTETIKE_kerkesat_qytetare_jan-gus_2026.xlsx")
    batch = []
    for i, r in enumerate(rows[3:], start=4):
        if not r[0] or not str(r[0]).startswith("KQ-"):
            continue
        created = dt.datetime.strptime(r[1], "%d.%m.%Y").date()
        closed = r[9].date() if r[9] else None
        batch.append(
            ("req", i, r[0], created, closed, r[4], r[5], c.normalize_admin_unit(r[6]),
             r[7], r[8], r[10])
        )  # fmt: skip
    insert_rows(con, "request", batch)

    # budget: thousands of lek
    rows = xlsx_rows("02_SINTETIKE_zbatimi_buxhetit_jan-gus_2026.xlsx")
    batch, plan_sum, actual_sum = [], 0.0, 0.0
    for i, r in enumerate(rows[3:], start=4):
        if r[0] is None:
            continue
        if r[0] == "TOTALI":
            totals["budget"] = {"plan": r[4], "actual": r[5]}
            continue
        month = c.parse_month(r[0])
        kind = c.normalize_line_type(r[3])
        batch.append(("bud", i, month, r[1], r[2], kind, r[4] * 1000, r[5] * 1000))
        plan_sum += r[4]
        actual_sum += r[5]
    insert_rows(con, "budget_line", batch)
    totals["budget_loaded"] = {"plan": plan_sum, "actual": actual_sum}

    # population (comma CSV, title row 1, header row 2)
    rows = csv_rows("ref_SINTETIKE_popullsia_njesite.csv", ",")
    batch = [("pop", i, r[0], r[1], int(r[2])) for i, r in enumerate(rows[2:], start=3)]
    insert_rows(con, "population", batch)

    # waste (';' CSV, decimal comma)
    rows = csv_rows("zarfi-1_SINTETIKE_pastrimi_mbetjet_2026.csv", ";")
    batch, tonnes_sum = [], 0.0
    for i, r in enumerate(rows[2:], start=3):
        if not r[0]:
            continue
        if r[0] == "TOTALI":
            totals["waste"] = {"tonnes": c.parse_number(r[2]), "trips": c.parse_number(r[3])}
            continue
        t = c.parse_number(r[2])
        batch.append(
            ("wst", i, c.parse_month(r[0]), c.normalize_admin_unit(r[1]), t,
             c.parse_int(r[3]), c.parse_int(r[4]))
        )  # fmt: skip
        tonnes_sum += t
    insert_rows(con, "waste_collection", batch)
    totals["waste_loaded"] = {"tonnes": tonnes_sum}

    # revenue: month cells are real dates; thousands of lek
    rows = xlsx_rows("zarfi-2_SINTETIKE_taksat_tarifat_arketimi_2026.xlsx")
    batch, coll_sum = [], 0.0
    for i, r in enumerate(rows[3:], start=4):
        if r[0] is None:
            continue
        if r[0] == "Gjithsej":
            totals["revenue"] = {"collected": r[4]}
            continue
        batch.append(("rev", i, c.parse_month(r[0]), r[1], r[2], r[3] * 1000, r[4] * 1000))
        coll_sum += r[4]
    insert_rows(con, "revenue", batch)
    totals["revenue_loaded"] = {"collected": coll_sum}

    # staff: as_of from the header date
    rows = xlsx_rows("zarfi-3_SINTETIKE_burimet_njerezore_2026.xlsx")
    as_of = c.find_date_in_text(rows[2][2])
    batch = []
    for i, r in enumerate(rows[3:], start=4):
        if r[1] is None:
            continue
        if r[1] == "Gjithsej":
            totals["staff"] = {"headcount": r[2]}
            continue
        batch.append(("hr", i, r[1], r[2], r[3], r[4], as_of))
    insert_rows(con, "staff", batch)
    return totals


@pytest.fixture
def loaded(db):
    totals = load_all(db)
    return db, totals


def test_total_rows_reconcile(loaded):
    _, totals = loaded
    assert totals["budget"]["plan"] == pytest.approx(totals["budget_loaded"]["plan"], abs=0.05)
    assert totals["budget"]["actual"] == pytest.approx(totals["budget_loaded"]["actual"], abs=0.05)
    assert totals["waste"]["tonnes"] == pytest.approx(totals["waste_loaded"]["tonnes"], abs=0.05)
    assert totals["revenue"]["collected"] == pytest.approx(
        totals["revenue_loaded"]["collected"], abs=0.05
    )


def test_population_is_synthetic_not_official(loaded):
    con, _ = loaded
    totals = dict(con.execute("SELECT basis, sum(residents) FROM population GROUP BY 1").fetchall())
    assert totals == {"census_2023": 121_400, "civil_registry": 208_700}


def test_all_13_passports_compute_on_samples(loaded):
    con, _ = loaded
    assert reg.coverage(con)["computable"] == 13
    for p in reg.load_pack().passports:
        res = reg.compute(p, con)
        assert res.value is not None, p.code
        assert res.period == "2026-08", p.code
        assert res.lineage, p.code


def test_planted_patterns(loaded):
    con, _ = loaded
    v = {p.code: reg.compute(p, con) for p in reg.load_pack().passports}
    assert v["REQ-01"].value == 2400
    assert 35 <= v["FIN-02"].value <= 40  # capital execution far from the 70% target
    assert reg.status_vs_target(reg.get_passport("FIN-02"), v["FIN-02"].value) == "off_track"
    assert 50 <= v["REV-02"].value <= 60  # cleaning-fee coverage
    cost = [s.value for s in v["WST-03"].series]
    assert max(abs(b / a - 1) for a, b in zip(cost, cost[1:], strict=False)) > 0.5
    census = v["WST-02"].value
    civil = reg.compute(reg.get_passport("WST-02"), con, "civil_registry").value
    assert census > civil * 1.5  # the basis choice matters
    # Gjinar reports 0 tonnes in July
    rows = con.execute(reg.get_passport("WST-01").breakdown_sql).fetchall()
    assert ("2026-07", "Gjinar", 0.0) in rows
    # public works on-time rate drops in July-August
    by_dept = defaultdict(dict)
    for period, dept, value in con.execute(reg.get_passport("REQ-02").breakdown_sql).fetchall():
        by_dept[dept][period] = value
    pw = by_dept["Drejtoria e Punëve Publike"]
    assert min(pw[f"2026-0{m}"] for m in range(1, 7)) > 80
    assert max(pw["2026-07"], pw["2026-08"]) < 72
    assert 1450 <= con.execute("SELECT sum(headcount) FROM staff").fetchone()[0] <= 1550
