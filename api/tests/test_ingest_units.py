"""Unit tests for the ingest building blocks: reader, layout, PII gate, profiling, mapping."""

import datetime as dt
import io

import pytest
from openpyxl import Workbook

from app.catalog import get_dataset
from app.ingest import mapper, pii, profile, reader
from app.ingest.errors import IngestError
from app.ingest.layout import analyse, total_kind
from app.ingest.profile import ColumnProfile


def xlsx_bytes(build) -> bytes:
    wb = Workbook()
    build(wb.active)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------------------------
# reader
# --------------------------------------------------------------------------------------------


def test_csv_cp1252_semicolon_decimal_comma():
    raw = "Muaji;Njësia administrative;Sasia (ton)\n2026-01;Shirgjani;1.234,5\n2026-02;Papër;12,5\n"
    table = reader.read_table(raw.encode("cp1252"), "mbetje.csv")
    assert table.encoding == "cp1252" and table.delimiter == ";"
    assert table.rows[1] == ["2026-01", "Shirgjani", "1.234,5"]
    assert table.rows[2][1] == "Papër"


def test_csv_comma_with_quoted_decimal_comma_and_bom():
    raw = '\ufeffMuaji,Sasia,Kurse\n2026-01,"1.234,5",10\n2026-02,"2,5",3\n'
    table = reader.read_table(raw.encode("utf-8"), "a.csv")
    assert table.encoding == "utf-8-sig" and table.delimiter == ","
    assert table.rows[0][0] == "Muaji" and table.rows[1] == ["2026-01", "1.234,5", "10"]


def test_xlsx_merged_cells_are_forward_filled_and_totals_excluded():
    def build(ws):
        ws["A1"] = "SINTETIKE · Titulli i raportit"
        ws.merge_cells("A1:C1")
        ws.append(["Muaji", "Njësia administrative", "Sasia (ton)"])
        ws.append(["Janar 2026", "Elbasan", 10.5])
        ws.append([None, "Papri", 2])
        ws.append([None, "Tregani", 3])
        ws.merge_cells("A3:A5")
        ws.append(["Shuma", None, 15.5])

    table = reader.read_table(xlsx_bytes(build), "test.xlsx")
    assert table.kind == "xlsx" and len(table.merged) == 2
    layout = analyse(table)
    assert layout.header_row == 2
    assert [r[0] for _, r in layout.data] == ["Janar 2026"] * 3
    assert [(e.row_no, e.reason) for e in layout.excluded] == [(1, "title"), (6, "total_row")]
    assert layout.grand_total.cells[2] == 15.5


def test_two_level_header_and_unit_from_group_cell():
    def build(ws):
        ws.append(["Titulli SINTETIKE"])
        ws.append(["Programi", "Plani (000 lekë)", None, "Muaji"])
        ws.merge_cells("B2:C2")
        ws.append([None, "Korrente", "Kapitale", None])
        ws.append(["Arsimi", 10, 2, "2026-01"])
        ws.append(["Sporti", 5, 1, "2026-01"])

    layout = analyse(reader.read_table(xlsx_bytes(build), "b.xlsx"))
    assert layout.header_rows == [2, 3]
    assert layout.headers[:3] == ["Programi", "Plani (000 lekë) Korrente", "Plani (000 lekë) Kapitale"]
    assert layout.col_multiplier[1:3] == [1000, 1000]
    assert len(layout.data) == 2


def test_duplicate_and_empty_headers_become_unique():
    table = reader.read_table(b"Plani;Plani;;X\n1;2;3;a\n4;5;6;b\n", "d.csv")
    assert analyse(table).headers == ["Plani", "Plani (2)", "Kolona 3", "X"]


@pytest.mark.parametrize(
    ("row", "kind"),
    [
        (["TOTALI", None, 5], "total_row"),
        ([None, "Gjithsej", 1501], "total_row"),
        (["Shuma", 3], "total_row"),
        (["Nëntotali Janar", 3], "subtotal"),
        (["Grand total", 3], "total_row"),
        (["Taksa e ndërtesës", 3], None),
        (["2026-01", "Totali i ri", 3], "total_row"),
    ],
)
def test_total_rows(row, kind):
    assert total_kind(row) == kind


def test_upload_checks():
    with pytest.raises(IngestError) as exc:
        reader.check_upload(b"x", "notes.txt")
    assert exc.value.status_code == 415
    with pytest.raises(IngestError) as exc:
        reader.check_upload(b"0" * (reader.MAX_BYTES + 1), "big.csv")
    assert exc.value.status_code == 413
    with pytest.raises(IngestError) as exc:
        reader.read_table(b"not a zip", "broken.xlsx")
    assert exc.value.code == "unreadable_file"


@pytest.mark.parametrize(
    ("raw", "clean"),
    [
        ("../../etc/passwd.csv", "passwd.csv"),
        ("C:\\Users\\x\\raport.xlsx", "raport.xlsx"),
        ("zarfi 2 <arkëtimi>.xlsx", "zarfi 2 _arkëtimi_.xlsx"),
        ("", "skedar"),
        ("a" * 200 + ".csv", "a" * 116 + ".csv"),
    ],
)
def test_sanitize_filename(raw, clean):
    assert reader.sanitize_filename(raw) == clean


def test_xls_reader():
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("S")
    for j, v in enumerate(["Muaji", "Njësia administrative", "Sasia (ton)"]):
        ws.write(0, j, v)
    ws.write(1, 0, "2026-01")
    ws.write(1, 1, "Gjinar")
    ws.write(1, 2, 4.5)
    buf = io.BytesIO()
    wb.save(buf)
    table = reader.read_table(buf.getvalue(), "old.xls")
    assert table.kind == "xls" and table.rows[1] == ["2026-01", "Gjinar", 4.5]


# --------------------------------------------------------------------------------------------
# PII gate
# --------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("header", "values", "kind"),
    [
        ("Emri i kërkuesit", ["Arben K.", "Ilir M."], "name"),
        ("Emri dhe mbiemri", ["x"], "name"),
        ("Nr. telefoni", ["+355 69 000 11 22"], "phone"),
        ("Cel.", ["x"], "phone"),
        ("E-mail", ["x"], "email"),
        ("Adresa", ["Rr. Kristoforidhi 5"], "address"),
        ("Nr. personal", ["x"], "personal_id"),
        ("Kodi", ["J12345678K", "K98765432L", "I11223344M"], "personal_id"),
        ("Kontakti", ["a.b@example.com", "c@d.al", "e@f.org"], "email"),
        ("Shënime", ["Telefononi 069 123 4567 për detaje të mëtejshme"], "phone"),
        ("Përgjegjësi", [f"{n} {i}." for n, i in zip("Arben Ilir Teuta Genc Sara".split(), "KLMNP")], "name"),
        ("Qytetari", ["Arben Hoxha", "Teuta Leka", "Ilir Gjoka", "Sara Bala", "Genc Duka"], "name"),
    ],
)
def test_pii_detected(header, values, kind):
    finding = pii.detect_column(header, values)
    assert finding is not None and finding.kind == kind


@pytest.mark.parametrize(
    ("header", "values"),
    [
        ("Emri i programit", ["Arsimi bazë", "Mbrojtja civile", "Sporti dhe argëtimi"]),
        ("Plani (lekë)", [680000000, 690123456, 670000000]),
        ("Operatori", ["Operatori A", "Operatori B"] * 20),
        ("Drejtoria", ["Policia Bashkiake", "Drejtoria e Financës", "Kabineti i Kryetarit"]),
        ("Kanali (telefon/online)", ["Telefon", "Online", "Sportel"]),
        ("Drejtoria përgjegjëse", ["Drejtoria e Punëve Publike"]),
        ("Nr.", ["KQ-2026-00001", "KQ-2026-00002"]),
        ("Kategoria", ["Rrugë dhe trotuare", "Ndriçimi publik"]),
        ("Kodi i programit", ["01110", "05100"]),
    ],
)
def test_pii_not_detected(header, values):
    assert pii.detect_column(header, values) is None


def test_mask_sample():
    assert pii.mask_sample("2526229.1") == "25•••29.1"
    assert pii.mask_sample("shkruani në a@b.al") == "shkruani në [email]"
    assert pii.mask_sample("J12345678K") == "[nr. personal]"
    assert pii.mask_sample("+355 69 000 11 22") == "[telefon]"
    assert pii.mask_sample("2026-01-01") == "2026-01-01"
    assert pii.mask_sample("KQ-2026-00001") == "KQ-2026-00001"


# --------------------------------------------------------------------------------------------
# profiling
# --------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("value", "kind", "fmt"),
    [
        ("29.01.2026", "date", "dmy"),
        ("Janar 2026", "date", "month_name"),
        ("2026-01", "date", "iso_month"),
        (dt.datetime(2026, 1, 1), "date", "excel_date"),
        ("2.562,4", "float", "decimal_comma"),
        ("23.944", "int", None),
        ("01110", "string", "code"),
        ("KQ-2026-00001", "string", None),
        ("2400 kërkesa", "string", None),
        (12, "int", None),
        (1.5, "float", None),
        ("Elbasan", "string", None),
    ],
)
def test_classify(value, kind, fmt):
    assert profile.classify(value) == (kind, fmt)


def test_profile_column_null_share_and_distinct_samples():
    values = ["A", "A", None, "B", "C", "D", "E", "F", None, None]
    col = profile.profile_column(0, "X", values, None)
    assert col.null_pct == 30.0 and col.samples == ["A", "B", "C", "D", "E"]
    assert col.inferred_type == "string"
    dropped = profile.profile_column(1, "Emri", ["Arben K."], pii.PiiFinding("name", "header", 1))
    assert dropped.samples == [] and dropped.dropped and dropped.pii == "name"
    assert profile.profile_column(2, "Y", [None, None], None).inferred_type == "empty"


# --------------------------------------------------------------------------------------------
# mapping
# --------------------------------------------------------------------------------------------


def cols(*pairs: tuple[str, str]) -> list[ColumnProfile]:
    return [
        ColumnProfile(i, name, typ, [], 0.0, None, False) for i, (name, typ) in enumerate(pairs)
    ]


def test_detect_dataset_from_headers():
    budget = cols(
        ("Muaji", "date"),
        ("Kodi i programit", "string"),
        ("Programi", "string"),
        ("Plani (000 lekë)", "float"),
        ("Fakti (000 lekë)", "float"),
    )
    ranked = mapper.detect_dataset(budget, "raport.xlsx")
    assert ranked[0]["key"] == "budget" and ranked[0]["score"] >= 0.8
    unknown = mapper.detect_dataset(cols(("Foo", "string"), ("Bar", "int")), "x.csv")
    assert unknown[0]["score"] < mapper.DATASET_THRESHOLD


def test_rules_mapping_is_one_to_one_and_capped():
    ds = get_dataset("budget")
    columns = cols(
        ("Programi", "string"),
        ("Kodi i programit", "string"),
        ("Muaji", "date"),
        ("Plani vjetor 2026", "float"),
        ("Realizimi", "float"),
        ("Shënime", "string"),
    )
    sugg, _ = mapper.rules_mapping(ds, columns)
    got = {s.column: s.field for s in sugg}
    assert got == {
        "Programi": "programme",
        "Kodi i programit": "programme_code",
        "Muaji": "month",
        "Plani vjetor 2026": "planned_lek",
        "Realizimi": "actual_lek",
        "Shënime": None,
    }
    assert all(s.confidence <= mapper.RULES_CAP for s in sugg)


def test_rules_respect_value_types():
    ds = get_dataset("waste")
    sugg, _ = mapper.rules_mapping(ds, cols(("Sasia", "string"), ("Muaji", "int")))
    assert {s.column: s.field for s in sugg} == {"Sasia": None, "Muaji": None}


def test_rules_question_for_a_missing_required_field():
    ds = get_dataset("waste")
    sugg, _ = mapper.rules_mapping(ds, cols(("Pesha totale", "float")))
    assert sugg[0].field == "tonnes"  # "Pesha" is a synonym contained in the header
    columns = cols(("Muaji", "date"), ("Njësia", "string"), ("Mbeturinat e mbledhura", "float"))
    sugg, cands = mapper.rules_mapping(ds, columns)
    q = mapper.rules_question(ds, columns, sugg, cands)
    assert q["column"] == "Mbeturinat e mbledhura"
    assert [o["field"] for o in q["options"]] == ["tonnes", None]
    assert q["text"]["sq"] and q["text"]["en"]


def test_synonym_keys_keep_semantic_parentheses():
    assert mapper.synonym_key("Programi (kodi)") == "programi kodi"
    assert mapper.synonym_key("Sasia (ton)") == "sasia"
    assert mapper.synonym_key("Plani (000 lekë)") == "plani"
