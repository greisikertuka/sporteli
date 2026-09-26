"""Report exports: the KPI workbook with visible provenance, and the open-data CSV.

The workbook's "Treguesit" sheet puts a visible "Burimi" (source) column next to every value
— file, rows, passport version, time of computation — and repeats that provenance in a cell
comment on the value. "Burimet" lists the source files with their SHA-256 hashes and load
counts; "Shënime" carries the draft-formula and synthetic-data disclaimers.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
from collections.abc import Sequence

import duckdb
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.catalog import OWNERS, POPULATION_BASES, get_dataset, t
from app.indicators import registry as reg
from app.indicators import signals as sig
from app.indicators.service import Evaluated, SourceMeta, source_meta

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MEDIA_TYPE = "text/csv; charset=utf-8"
LICENSE_URL = "https://creativecommons.org/licenses/by/4.0/"

SHEETS = {
    "indicators": t("Treguesit", "Indicators"),
    "sources": t("Burimet", "Sources"),
    "notes": t("Shënime", "Notes"),
}

INDICATOR_COLUMNS: dict[str, list[tuple[str, int]]] = {
    "sq": [
        ("Kodi", 9),
        ("Treguesi", 42),
        ("Vlera", 14),
        ("Njësia", 18),
        ("Periudha", 13),
        ("Statusi", 26),
        ("Objektivi", 12),
        ("Burimi", 78),
        ("Baza", 16),
        ("Sinjalet", 70),
    ],
    "en": [
        ("Code", 9),
        ("Indicator", 42),
        ("Value", 14),
        ("Unit", 18),
        ("Period", 13),
        ("Status", 26),
        ("Target", 12),
        ("Source", 78),
        ("Basis", 16),
        ("Signals", 70),
    ],
}

SOURCE_COLUMNS: dict[str, list[tuple[str, int]]] = {
    "sq": [
        ("Skedari", 52),
        ("Të dhënat", 30),
        ("SHA-256", 30),
        ("Rreshta të lexuar", 12),
        ("Rreshta të ngarkuar", 12),
        ("Rreshta të përjashtuar", 12),
        ("Kontrolli i totalit", 30),
        ("Ngarkuar më", 18),
        ("Sintetike", 10),
        ("Treguesit", 40),
    ],
    "en": [
        ("File", 52),
        ("Dataset", 30),
        ("SHA-256", 30),
        ("Rows read", 12),
        ("Rows loaded", 12),
        ("Rows excluded", 12),
        ("Total-row check", 30),
        ("Loaded at", 18),
        ("Synthetic", 10),
        ("Indicators", 40),
    ],
}

STATUS_TEXT = {
    "on_track": t("Në objektiv", "On track"),
    "off_track": t("Jashtë objektivit", "Off track"),
    "no_target": t("Pa objektiv", "No target"),
    None: t("Pa vlerë", "No value"),
}
FORMULA_STATUS_TEXT = {
    "draft": t("draft — në pritje të validimit", "draft — awaiting validation"),
    "from_source": t("nga burimi zyrtar", "from the official source"),
    "validated": t("e validuar", "validated"),
}

NUMBER_FORMATS = {
    "count": "#,##0",
    "percent": "0.0",
    "days": "0.0",
    "kg_per_resident": "#,##0.0",
    "lek_per_ton": "#,##0",
    "per_1000": "0.0",
}

INK = "1F3A5F"
HEADER_FILL = PatternFill("solid", fgColor=INK)
HEADER_FONT = Font(bold=True, color="FFFFFF")
GRID = Side(style="thin", color="D0D7DE")
BORDER = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)
STATUS_FILLS = {
    "on_track": PatternFill("solid", fgColor="E3F1E7"),
    "off_track": PatternFill("solid", fgColor="FDF0DB"),
    "missing": PatternFill("solid", fgColor="F2F2F2"),
}
SYNTHETIC_FILL = PatternFill("solid", fgColor="FFF4CE")
MISSING_FONT = Font(italic=True, color="6B7280")
SOURCE_FONT = Font(color=INK)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


# --------------------------------------------------------------------------------------------
# Text helpers
# --------------------------------------------------------------------------------------------


def _computed_at(e: Evaluated) -> dt.datetime | None:
    if not e.result or not e.result.computed_at:
        return None
    try:
        return dt.datetime.fromisoformat(e.result.computed_at)
    except ValueError:
        return None


def _stamp(when: dt.datetime | None, loc: str) -> str:
    if when is None:
        return "—"
    when = when.astimezone(dt.UTC) if when.tzinfo else when
    fmt = "%d.%m.%Y %H:%M" if loc == "sq" else "%Y-%m-%d %H:%M"
    return f"{when.strftime(fmt)} UTC"


def short_ranges(ranges: str, row_count: int, loc: str, keep: int = 3) -> str:
    """ "4–7, 9, 11–15, 17" → "4–7, 9, 11–15 … (N rreshta)" when there are many segments."""
    parts = [x for x in ranges.split(", ") if x]
    if len(parts) <= keep:
        return ranges
    word = "rreshta" if loc == "sq" else "rows"
    return f"{', '.join(parts[:keep])} … ({reg.format_number(row_count, 0, loc)} {word})"


def provenance(e: Evaluated, loc: str = "sq") -> str:
    """The "Burimi" cell: "file · rreshtat … · pasaporta vX · llogaritur më …"."""
    if not e.result:
        return "—"
    rows_word = "rreshtat" if loc == "sq" else "rows"
    files = " + ".join(
        f"{s['filename']} · {rows_word} {short_ranges(s['row_ranges'], s['row_count'], loc)}"
        for s in e.lineage
    )
    passport = (
        f"pasaporta v{e.passport.version}" if loc == "sq" else f"passport v{e.passport.version}"
    )
    when = _stamp(_computed_at(e), loc)
    computed = f"llogaritur më {when}" if loc == "sq" else f"computed {when}"
    return f"{files or '—'} · {passport} · {computed}"


def provenance_comment(e: Evaluated, loc: str = "sq", max_ranges: int = 400) -> str:
    """Longer provenance for the value cell's comment (full ranges up to ``max_ranges``)."""
    p = e.passport
    lines = ["Burimi:" if loc == "sq" else "Source:"]
    for s in e.lineage:
        ranges = s["row_ranges"]
        if len(ranges) > max_ranges:
            ranges = ranges[:max_ranges].rsplit(", ", 1)[0] + " …"
        n = reg.format_number(s["row_count"], 0, loc)
        lines.append(f"• {s['filename']}")
        lines.append(
            f"  {'Rreshtat' if loc == 'sq' else 'Rows'}: {ranges} "
            f"({n} {'rreshta' if loc == 'sq' else 'rows'})"
        )
        if s["file_hash"]:
            lines.append(f"  SHA-256: {s['file_hash']}")
    status = FORMULA_STATUS_TEXT[p.formula_status][loc]
    label = "Pasaporta" if loc == "sq" else "Passport"
    lines.append(f"{label}: {p.code} v{p.version} · formula {status}")
    when = _stamp(_computed_at(e), loc)
    lines.append(f"{'Llogaritur më' if loc == 'sq' else 'Computed'}: {when}")
    if e.basis:
        basis_label = "Baza e popullsisë" if loc == "sq" else "Population basis"
        lines.append(f"{basis_label}: {POPULATION_BASES[e.basis][loc]}")
    if any(s["synthetic"] for s in e.sources):
        lines.append("TË DHËNA SINTETIKE" if loc == "sq" else "SYNTHETIC DATA")
    return "\n".join(lines)


def target_text(p: reg.Passport, loc: str) -> str:
    if p.target is None or p.direction == "none":
        return "—"
    sign = "≥" if p.direction == "higher_better" else "≤"
    value = reg.format_value(p.target, p.unit, loc, 0)
    unit = "" if p.unit == "percent" else f" {p.label[loc]}"
    return f"{sign} {value}{unit}"


def unit_text(p: reg.Passport, loc: str) -> str:
    return p.label[loc]


def status_text(e: Evaluated, loc: str) -> str:
    if e.state == "missing":
        parts = []
        for key in e.missing:
            ds = get_dataset(key)
            parts.append(
                f"Mungon: {ds.name['sq']} · Përgjegjës: {OWNERS[ds.owner]['sq']}"
                if loc == "sq"
                else f"Missing: {ds.name['en']} · Owner: {OWNERS[ds.owner]['en']}"
            )
        return "\n".join(parts)
    return STATUS_TEXT[reg.status_vs_target(e.passport, e.value)][loc]


def signals_text(e: Evaluated, loc: str) -> str:
    return "\n".join(f"• {s.message[loc]}" for s in e.signals) or "—"


def as_of(evaluated: Sequence[Evaluated]) -> str | None:
    periods = [e.period for e in evaluated if e.period]
    return max(periods) if periods else None


def xlsx_filename(evaluated: Sequence[Evaluated], loc: str = "sq") -> str:
    stem = "sportel-treguesit" if loc == "sq" else "sportel-indicators"
    return f"{stem}-{as_of(evaluated) or 'pa-te-dhena'}.xlsx"


# --------------------------------------------------------------------------------------------
# Workbook
# --------------------------------------------------------------------------------------------


def _header(ws: Worksheet, columns: list[tuple[str, int]]) -> None:
    for i, (title, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30


def _excluded_count(m: SourceMeta) -> int | None:
    if m.rows_excluded_json:
        try:
            data = json.loads(m.rows_excluded_json)
        except ValueError:
            data = None
        if isinstance(data, list):
            return sum(int(x.get("count", 0) or 0) for x in data if isinstance(x, dict))
        if isinstance(data, dict):
            return sum(int(v or 0) for v in data.values() if isinstance(v, int | float))
    if m.rows_read is not None and m.rows_loaded is not None:
        return max(m.rows_read - m.rows_loaded, 0)
    return None


def _reconciliation_text(m: SourceMeta, loc: str) -> str:
    items = sig.parse_reconciliation(m.filename, m.reconciliation_json)
    items = [r for r in items if r.file_total is not None]
    if not items:
        return "—"
    if any(r.stale for r in items):
        return (
            "↻ pjesërisht e zëvendësuar nga një ngarkim i mëvonshëm"
            if loc == "sq"
            else "↻ partly replaced by a later load"
        )
    bad = [r for r in items if not r.ok]
    if not bad:
        labels = ", ".join(r.label.get(loc) or r.field for r in items)
        return f"✓ përputhet ({labels})" if loc == "sq" else f"✓ matches ({labels})"
    labels = ", ".join(r.label.get(loc) or r.field for r in bad)
    return f"✗ nuk përputhet: {labels}" if loc == "sq" else f"✗ does not match: {labels}"


def _indicator_sheet(ws: Worksheet, evaluated: Sequence[Evaluated], loc: str) -> None:
    _header(ws, INDICATOR_COLUMNS[loc])
    for row, e in enumerate(evaluated, start=2):
        p = e.passport
        values = [
            p.code,
            p.name[loc],
            e.value if e.value is not None else "—",
            unit_text(p, loc),
            reg.format_period(p.reported_period(e.period), loc) if e.period else "—",
            status_text(e, loc),
            target_text(p, loc),
            provenance(e, loc),
            POPULATION_BASES[e.basis][loc] if e.basis else "—",
            signals_text(e, loc),
        ]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border = BORDER
            cell.alignment = WRAP_TOP if col in (2, 6, 8, 10) else TOP
        value_cell = ws.cell(row=row, column=3)
        if e.value is not None:
            value_cell.number_format = NUMBER_FORMATS.get(p.unit, "0.0")
            if p.decimals is not None:
                value_cell.number_format = "#,##0" + ("." + "0" * p.decimals if p.decimals else "")
            comment = Comment(provenance_comment(e, loc), "Sportel")
            comment.width = 420
            comment.height = 220
            value_cell.comment = comment
        else:
            value_cell.alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=row, column=8).font = SOURCE_FONT
        status_cell = ws.cell(row=row, column=6)
        if e.state == "missing":
            for col in range(1, len(values) + 1):
                ws.cell(row=row, column=col).font = MISSING_FONT
            status_cell.fill = STATUS_FILLS["missing"]
        else:
            fill = STATUS_FILLS.get(reg.status_vs_target(p, e.value) or "")
            if fill is not None:
                status_cell.fill = fill
    last_col = get_column_letter(len(INDICATOR_COLUMNS[loc]))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{last_col}{len(evaluated) + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def _sources_sheet(
    ws: Worksheet,
    evaluated: Sequence[Evaluated],
    meta: dict[str, SourceMeta],
    loc: str,
) -> None:
    _header(ws, SOURCE_COLUMNS[loc])
    used: dict[str, list[str]] = {}
    for e in evaluated:
        for s in e.sources:
            used.setdefault(s["source_id"], []).append(e.passport.code)
    yes, no = ("Po", "Jo") if loc == "sq" else ("Yes", "No")
    for row, m in enumerate(meta.values(), start=2):
        try:
            ds_name = get_dataset(m.dataset).name[loc] if m.dataset else "—"
        except KeyError:
            ds_name = m.dataset or "—"
        values = [
            m.filename,
            ds_name,
            m.file_hash or "—",
            m.rows_read if m.rows_read is not None else "—",
            m.rows_loaded if m.rows_loaded is not None else "—",
            _excluded_count(m) if _excluded_count(m) is not None else "—",
            _reconciliation_text(m, loc),
            _stamp(m.loaded_at, loc) if m.loaded_at else "—",
            yes if m.synthetic else no,
            ", ".join(used.get(m.id, [])) or "—",
        ]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border = BORDER
            cell.alignment = WRAP_TOP if col in (1, 3, 7, 10) else TOP
        if m.synthetic:
            ws.cell(row=row, column=9).fill = SYNTHETIC_FILL
    ws.freeze_panes = "A2"
    if meta:
        last_col = get_column_letter(len(SOURCE_COLUMNS[loc]))
        ws.auto_filter.ref = f"A1:{last_col}{len(meta) + 1}"


def _notes_sheet(
    ws: Worksheet,
    evaluated: Sequence[Evaluated],
    meta: dict[str, SourceMeta],
    basis: str,
    generated: dt.datetime,
    loc: str,
) -> None:
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 80
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 22
    ws.column_dimensions["F"].width = 36
    synthetic = any(m.synthetic for m in meta.values())
    period = reg.format_period(as_of(evaluated), loc) if as_of(evaluated) else "—"
    basis_label = POPULATION_BASES[basis][loc]
    # the stock (point-in-time) indicators come from the registry, never a hand-kept list
    stock = ", ".join(
        dict.fromkeys(e.passport.code for e in evaluated if e.passport.period_kind == "point")
    )
    if loc == "sq":
        lines = [
            ("Sportel — Raporto një herë, provo çdo numër", "title"),
            (
                f"Raporti i treguesve kryesorë · gjendja: {period} · "
                f"gjeneruar më {_stamp(generated, loc)}",
                None,
            ),
            ("", None),
        ]
        if synthetic:
            lines.append(
                (
                    "TË DHËNA SINTETIKE. Skedarët burimorë janë eksporte sintetike departamentesh, "
                    "të krijuara për demonstrim; nuk janë të dhëna zyrtare të bashkisë.",
                    "synthetic",
                )
            )
        lines += [
            (
                "Formulat janë draft — në pritje të validimit nga bashkia dhe AMVV. Statusi i "
                "secilës formulë jepet në tabelën më poshtë.",
                "bold",
            ),
            (
                "Çdo numër llogaritet nga kodi me SQL të fiksuar në pasaportën e treguesit. "
                "AI nuk shkruan asnjë shifër.",
                None,
            ),
            (
                "Kolona «Burimi» dhe komentet në qelizat e vlerave tregojnë skedarin, rreshtat, "
                "versionin e pasaportës dhe kohën e llogaritjes. Fleta «Burimet» liston "
                "skedarët me hash SHA-256 dhe kontrollin e totalit.",
                None,
            ),
            (
                f"Treguesit për banor përdorin bazën e popullsisë: {basis_label}.",
                None,
            ),
            (
                "Vlerat janë nga 1 janari deri në fund të muajit të fundit me të dhëna"
                + (
                    f", përveç treguesve të gjendjes në fund të muajit ({stock})." if stock else "."
                ),
                None,
            ),
            (
                "Përgjegjësit janë role të përkohshme; emrat konfirmohen me bashkinë.",
                None,
            ),
            ("", None),
        ]
        table_head = [
            "Kodi",
            "Treguesi",
            "Formula",
            "Statusi i formulës",
            "Referenca SMP",
            "Përgjegjës",
        ]
    else:
        lines = [
            ("Sportel — Report once, prove every number", "title"),
            (
                f"Core indicator report · as of {period} · generated {_stamp(generated, loc)}",
                None,
            ),
            ("", None),
        ]
        if synthetic:
            lines.append(
                (
                    "SYNTHETIC DATA. The source files are synthetic department exports created "
                    "for demonstration; they are not official municipal data.",
                    "synthetic",
                )
            )
        lines += [
            (
                "Formulas are drafts awaiting validation by the municipality and AMVV. The "
                "status of each formula is given in the table below.",
                "bold",
            ),
            (
                "Every number is computed by code with the fixed SQL of the indicator passport. "
                "AI does not type a single digit.",
                None,
            ),
            (
                "The “Source” column and the comments on value cells give the file, the rows, "
                "the passport version and the time of computation. The “Sources” sheet lists "
                "the files with their SHA-256 hash and the total-row check.",
                None,
            ),
            (f"Per-resident indicators use the population basis: {basis_label}.", None),
            (
                "Values run from 1 January to the end of the latest month with data"
                + (
                    f", except stock indicators at the end of the month ({stock})."
                    if stock
                    else "."
                ),
                None,
            ),
            ("Owners are placeholder roles; names are confirmed with the municipality.", None),
            ("", None),
        ]
        table_head = ["Code", "Indicator", "Formula", "Formula status", "SMP reference", "Owner"]
    row = 1
    for text, style in lines:
        cell = ws.cell(row=row, column=1, value=text or None)
        if text:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if style == "title":
            cell.font = Font(bold=True, size=14, color=INK)
        elif style == "bold":
            cell.font = Font(bold=True)
        elif style == "synthetic":
            cell.font = Font(bold=True, color="7A4B00")
            cell.fill = SYNTHETIC_FILL
        if text and len(text) > 120:
            ws.row_dimensions[row].height = 30
        row += 1
    for col, title in enumerate(table_head, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for e in evaluated:
        row += 1
        p = e.passport
        ref = p.smp_ref or "—"
        if p.smp_note:
            ref = f"{ref} · {p.smp_note[loc]}"
        values = [
            p.code,
            p.name[loc],
            p.formula[loc],
            FORMULA_STATUS_TEXT[p.formula_status][loc],
            ref,
            OWNERS[p.owner][loc],
        ]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border = BORDER
            cell.alignment = WRAP_TOP


def build_xlsx(
    evaluated: Sequence[Evaluated],
    con: duckdb.DuckDBPyConnection,
    loc: str = "sq",
    generated: dt.datetime | None = None,
) -> bytes:
    """The KPI workbook as bytes (sheets Treguesit / Burimet / Shënime, or English names)."""
    loc = loc if loc in ("sq", "en") else "sq"
    generated = generated or dt.datetime.now(dt.UTC)
    meta = source_meta(con)
    basis = reg.get_population_basis(con)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEETS["indicators"][loc]
    _indicator_sheet(ws, evaluated, loc)
    _sources_sheet(wb.create_sheet(SHEETS["sources"][loc]), evaluated, meta, loc)
    _notes_sheet(wb.create_sheet(SHEETS["notes"][loc]), evaluated, meta, basis, generated, loc)
    wb.properties.creator = "Sportel"
    wb.properties.title = (
        f"Sportel — {'Treguesit kryesorë' if loc == 'sq' else 'Core indicators'}"
        f" {as_of(evaluated) or ''}".strip()
    )
    wb.properties.subject = "Raporto një herë, provo çdo numër"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------------------------
# Open data CSV
# --------------------------------------------------------------------------------------------

OPEN_DATA_COLUMNS = [
    "code",
    "name_sq",
    "name_en",
    "value",
    "unit",
    "period",
    "basis",
    "sources",
    "version",
    "formula_status",
    "target",
    "status",
    "synthetic",
    "computed_at",
]


def _plain_number(value: float | None) -> str:
    if value is None:
        return ""
    rounded = round(value, 4)
    if float(rounded).is_integer():
        return str(int(rounded))
    return repr(rounded)


def build_open_data_csv(evaluated: Sequence[Evaluated]) -> bytes:
    """Computable indicators as UTF-8 CSV with BOM (comma-separated, CRLF, header row)."""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow(OPEN_DATA_COLUMNS)
    for e in evaluated:
        if e.state != "computable":
            continue
        p = e.passport
        sources = " | ".join(
            f"{s['filename']} (sha256:{s['file_hash'][:12]})" if s["file_hash"] else s["filename"]
            for s in e.lineage
        )
        writer.writerow(
            [
                p.code,
                p.name["sq"],
                p.name["en"],
                _plain_number(e.value),
                p.unit,
                p.reported_period(e.period) or "",
                e.basis or "",
                sources,
                p.version,
                p.formula_status,
                _plain_number(p.target),
                reg.status_vs_target(p, e.value) or "",
                "true" if any(s["synthetic"] for s in e.sources) else "false",
                e.result.computed_at if e.result else "",
            ]
        )
    return buf.getvalue().encode("utf-8-sig")
